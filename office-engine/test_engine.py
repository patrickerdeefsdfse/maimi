"""Real-format regression fixtures; no employee documents or corporate calls."""
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from zipfile import ZipFile, ZIP_DEFLATED

ROOT = Path(os.environ.get('OFFICE_TEST_ENGINE_ROOT', str(Path(__file__).resolve().parent)))
spec = importlib.util.spec_from_file_location('office_engine', ROOT/'engine.py')
engine = importlib.util.module_from_spec(spec)
spec.loader.exec_module(engine)


class Office(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='megmeet-office-test-')
        self.base = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def call(self, operation, **kwargs):
        return engine.execute({'protocol':engine.PROTOCOL,'workspace':str(self.base),'operation':operation,**kwargs})

    def test_word_roundtrip_and_runs_tables(self):
        self.call('create',output='中文报告.docx',options={'paragraphs':[{'text':'办公报告','heading':1},'old value'], 'tables':[[['类别','数量'],['old',12]]]})
        original = (self.base/'中文报告.docx').read_bytes()
        from docx import Document
        doc = Document(self.base/'中文报告.docx')
        paragraph = doc.add_paragraph()
        paragraph.add_run('cro').bold = True
        paragraph.add_run('ss-run')
        doc.save(self.base/'runs.docx')
        self.call('edit',input='runs.docx',output='edited.docx',options={'replacements':[{'find':'old','replace':'new'},{'find':'cross-run','replace':'changed'}],'table_cells':[{'table':0,'row':1,'column':1,'text':'25'}]})
        value = self.call('read',input='edited.docx')
        self.assertIn('new value',value['text'])
        self.assertIn('changed',value['text'])
        self.assertIn('25',value['text'])
        self.assertEqual((self.base/'中文报告.docx').read_bytes(),original)
        self.assertTrue(Document(self.base/'edited.docx').paragraphs[-1].runs[0].bold)

    def test_excel_values_formulas_styles_chart(self):
        self.call('create',output='表格.xlsx',options={'sheets':[{'name':'销售','rows':[['部门','金额'],['A',12],['B',25]],'cells':[{'cell':'B4','formula':'=SUM(B2:B3)','number_format':'0.00','bold':True,'fill':'FFFF00'}],'freeze_panes':'A2','auto_filter':True,'column_widths':{'A':22},'charts':[{'title':'金额','min_col':2,'max_col':2,'max_row':3}]}]})
        result = self.call('read',input='表格.xlsx',options={'sheet':'销售','max_rows':3})
        self.assertTrue(result['more_rows'])
        self.assertEqual(result['rows'][1][1],12)
        import openpyxl
        book = openpyxl.load_workbook(self.base/'表格.xlsx')
        self.assertEqual(book['销售']['B4'].value,'=SUM(B2:B3)')
        self.assertTrue(book['销售']['B4'].font.bold)
        self.assertEqual(len(book['销售']._charts),1)
        book.close()
        self.call('edit',input='表格.xlsx',output='表格改.xlsx',options={'sheets':[{'name':'销售','cells':[{'cell':'B2','value':50}]}]})
        self.assertEqual(self.call('read',input='表格改.xlsx')['rows'][1][1],50)

    def test_excel_formula_injection_not_implicit(self):
        self.call('create',output='text.xlsx',options={'sheets':[{'name':'数据','rows':[['=1+1']], 'cells':[{'cell':'B1','value':'=WEBSERVICE("https://example.com")'}]}]})
        import openpyxl
        workbook = openpyxl.load_workbook(self.base/'text.xlsx')
        self.assertEqual(workbook['数据']['A1'].data_type,'s')
        self.assertEqual(workbook['数据']['B1'].data_type,'s')
        with self.assertRaises(engine.OfficeError):
            self.call('create',output='blocked.xlsx',options={'sheets':[{'name':'数据','cells':[{'cell':'B1','formula':'=WEBSERVICE("https://example.com")'}]}]})

    def test_streaming_excel_without_dimensions(self):
        import openpyxl
        book = openpyxl.Workbook(write_only=True)
        sheet = book.create_sheet('Streaming')
        sheet.append(['A','B'])
        sheet.append([10,20])
        book.save(self.base/'stream.xlsx')
        value = self.call('read',input='stream.xlsx',options={'max_columns':2,'max_rows':1})
        self.assertEqual(value['rows'],[['A','B']])
        self.assertTrue(value['more_rows'])
        value = self.call('read',input='stream.xlsx',options={'max_columns':2,'start_row':2})
        self.assertEqual(value['rows'],[[10,20]])

    def test_presentation_roundtrip(self):
        self.call('create',output='slides.pptx',options={'slides':[{'title':'Title old','bullets':['One','Two'],'notes':'备注'}]})
        self.call('edit',input='slides.pptx',output='slides-new.pptx',options={'replacements':[{'find':'old','replace':'new'}],'slides':[{'title':'Second','bullets':['Three']}]})
        result = self.call('read',input='slides-new.pptx')
        self.assertEqual(result['slide_count'],2)
        self.assertIn('Title new',result['text'])
        self.assertIn('备注',result['text'])

    def test_pdf_create_read_merge_rotate_extract_watermark(self):
        self.call('create',output='a.pdf',options={'paragraphs':['MEGMEET OFFICE','中文办公测试']})
        parsed = self.call('read',input='a.pdf')
        self.assertIn('MEGMEET',parsed['text'])
        self.call('pdf',output='merge.pdf',options={'action':'merge','inputs':['a.pdf','a.pdf']})
        self.assertEqual(self.call('inspect',input='merge.pdf')['page_count'],2)
        self.call('pdf',input='merge.pdf',output='rotate.pdf',options={'action':'rotate','pages':[2],'degrees':90})
        from pypdf import PdfReader
        self.assertEqual(PdfReader(self.base/'rotate.pdf').pages[0].rotation,90)
        self.call('pdf',input='a.pdf',output='water.pdf',options={'action':'watermark','text':'内部资料'})
        self.assertIn('内部资料',self.call('read',input='water.pdf')['text'])

    def test_pdf_forms(self):
        from reportlab.pdfgen.canvas import Canvas
        file = self.base/'form.pdf'
        c = Canvas(str(file))
        c.acroForm.textfield(name='employee',x=30,y=600,width=100,height=20)
        c.showPage()
        c.save()
        self.call('pdf',input='form.pdf',output='filled.pdf',options={'action':'fill_forms','fields':{'employee':'M000001'}})
        from pypdf import PdfReader
        self.assertEqual(PdfReader(self.base/'filled.pdf').get_fields()['employee']['/V'],'M000001')

    def test_image_transform(self):
        from PIL import Image
        Image.new('RGBA',(200,100),'orange').save(self.base/'source.png')
        self.call('image',input='source.png',output='new.jpg',options={'width':100,'rotate':90})
        result = self.call('inspect',input='new.jpg')
        self.assertEqual((result['width'],result['height']),(50,100))

    def test_csv_json_unicode(self):
        self.call('create',output='out.csv',options={'rows':[['名称','数量'],['测试','1'],['=2+2','x']]})
        self.assertIn("'=2+2",self.call('read',input='out.csv')['text'])
        self.call('create',output='out.json',options={'text':'{"中文":true}'})
        self.assertIn('中文',self.call('read',input='out.json')['text'])
        self.call('create',output='numbers.csv',options={'rows':[[-25,25,None,'=2+2']]})
        self.assertEqual((self.base/'numbers.csv').read_text(encoding='utf-8-sig').strip(), "-25,25,,'=2+2")

    def test_output_exists_and_traversal(self):
        (self.base/'keep.txt').write_text('untouched')
        for name in ('keep.txt','../escape.txt','/tmp/escape.txt','C:\\escape.txt','bad:stream.txt','CON.txt','x/../escape.txt'):
            with self.subTest(name=name), self.assertRaises(engine.OfficeError):
                self.call('create',output=name,options={'text':'changed'})
        self.assertEqual((self.base/'keep.txt').read_text(),'untouched')

    def test_output_parent_symlink_and_merge_escape(self):
        with tempfile.TemporaryDirectory() as external:
            try:
                (self.base/'linked').symlink_to(external,target_is_directory=True)
            except OSError:
                pass
            else:
                with self.assertRaises(engine.OfficeError):
                    self.call('create',output='linked/out.txt',options={'text':'secret'})
            with self.assertRaises(engine.OfficeError):
                self.call('pdf',output='merged.pdf',options={'action':'merge','inputs':['../other.pdf']})

    def test_archive_entities_rejected(self):
        with ZipFile(self.base/'bad.docx','w',ZIP_DEFLATED) as z:
            z.writestr('word/document.xml','<!DOCTYPE x [<!ENTITY e SYSTEM "file:///etc/passwd">]><x>&e;</x>')
        with self.assertRaises(engine.OfficeError) as caught:
            self.call('read',input='bad.docx')
        self.assertEqual(caught.exception.code,'UNSAFE_XML')

    def test_cli_protocol_and_private_error(self):
        request = {'protocol':engine.PROTOCOL,'operation':'read','workspace':str(self.base),'input':'missing.docx'}
        result = subprocess.run([sys.executable,'-I',str(ROOT/'engine.py')],input=json.dumps(request).encode(),capture_output=True,check=True)
        value = json.loads(result.stdout)
        self.assertFalse(value['ok'])
        self.assertNotIn(str(self.base),result.stdout.decode())
        request.update(operation='capabilities')
        result = subprocess.run([sys.executable,'-I',str(ROOT/'engine.py')],input=json.dumps(request).encode(),capture_output=True,check=True)
        self.assertEqual(json.loads(result.stdout)['result']['version'],'1.0.0')

    @unittest.skipUnless((ROOT/'ocr-config.json').exists(), 'OCR setup not complete')
    def test_offline_ocr(self):
        from PIL import Image, ImageDraw, ImageFont
        image = Image.new('RGB',(800,180),'white')
        font = ImageFont.truetype('C:/Windows/Fonts/arial.ttf',48) if sys.platform == 'win32' else ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',48)
        ImageDraw.Draw(image).text((20,50),'MEGMEET OFFICE 123',font=font,fill='black')
        image.save(self.base/'ocr.png')
        # Test exactly the CLI worker, whose audit hook refuses network access.
        request = {'protocol':engine.PROTOCOL,'operation':'read','workspace':str(self.base),'input':'ocr.png'}
        process = subprocess.run([sys.executable,'-I',str(ROOT/'engine.py')],input=json.dumps(request).encode(),capture_output=True,timeout=100)
        result = json.loads(process.stdout)
        self.assertTrue(result['ok'],result)
        self.assertIn('123',result['result']['text'])

    @unittest.skipUnless(engine.libreoffice(), 'LibreOffice setup not complete')
    def test_offline_convert_legacy_recalculate_pdf(self):
        self.call('create',output='source.docx',options={'paragraphs':['OFFICE LEGACY TEST']})
        self.call('convert',input='source.docx',output='old.doc')
        self.assertIn('OFFICE LEGACY TEST',self.call('read',input='old.doc')['text'])
        self.call('convert',input='source.docx',output='export.pdf')
        self.assertIn('OFFICE',self.call('read',input='export.pdf')['text'])
        self.call('create',output='calc.xlsx',options={'sheets':[{'name':'Data','rows':[['Value'],[12],[13]],'cells':[{'cell':'A4','formula':'=SUM(A2:A3)'}]}]})
        self.call('convert',input='calc.xlsx',output='recalc.xlsx')
        import openpyxl
        book = openpyxl.load_workbook(self.base/'recalc.xlsx',data_only=True)
        self.assertEqual(book['Data']['A4'].value,25)


if __name__ == '__main__':
    unittest.main(verbosity=2)
