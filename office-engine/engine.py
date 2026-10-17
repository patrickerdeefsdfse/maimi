"""MEGMEET Office v1. DSH-independent, bounded JSON worker; never edits in place.

Invoke with an isolated Python. The adapter supplies a trusted workspace, stages
only current-session attachments, enforces session policy, and kills on timeout.
No request can choose an executable or load a plugin; macros are disabled.
"""
from __future__ import annotations
import contextlib
import csv
import hashlib
import io
import itertools
import json
import os
from pathlib import Path
import re
import shutil
import signal
import subprocess
import sys
import tempfile
import warnings
from zipfile import ZipFile, is_zipfile

PROTOCOL = 'megmeet-office-v1'
VERSION = '1.0.0'
ROOT = Path(__file__).resolve().parent
MAX_INPUT = 64 * 1024 * 1024
MAX_OUTPUT = 128 * 1024 * 1024
MAX_TEXT = 120000
IMAGE_EXT = {'.png', '.jpg', '.jpeg', '.webp', '.tif', '.tiff', '.bmp'}
TEXT_EXT = {'.txt', '.md', '.csv', '.tsv', '.json', '.html', '.htm', '.xml', '.rtf'}
LEGACY = {'.doc': 'docx', '.ppt': 'pptx', '.odt': 'docx', '.ods': 'xlsx', '.odp': 'pptx'}
READABLE = IMAGE_EXT | TEXT_EXT | set(LEGACY) | {'.docx', '.xlsx', '.xls', '.pptx', '.pdf'}


class OfficeError(Exception):
    def __init__(self, code, message):
        self.code, self.message = code, message
        super().__init__(message)


def require(value, code='INVALID_REQUEST', message='Invalid office operation or parameters.'):
    if not value:
        raise OfficeError(code, message)


def integer(value, default, minimum, maximum):
    if value is None:
        return default
    require(type(value) is int and minimum <= value <= maximum)
    return value


def string(value, maximum=MAX_TEXT):
    require(isinstance(value, str) and len(value) <= maximum)
    return value


def workspace_path(base: Path, value: str, *, output=False):
    string(value, 1024)
    # Identical path contract on Windows and Linux: no URI, ADS, UNC or device.
    require(value and not re.search(r'[\x00-\x1f:]', value) and not value.startswith(('/', '\\')),
            'PATH_DENIED', 'Use a relative path inside the current workspace.')
    parts = value.replace('\\', '/').split('/')
    require(all(p not in ('', '.', '..') and not p.endswith((' ', '.'))
                and not re.match(r'(?i)^(con|prn|aux|nul|com[1-9]|lpt[1-9])(?:\.|$)', p)
                for p in parts), 'PATH_DENIED', 'Unsafe file path.')
    result = base.joinpath(*parts)
    for candidate in [result, *result.parents]:
        if candidate == base:
            break
        require(not candidate.is_symlink() and not (hasattr(candidate, 'is_junction') and candidate.is_junction()),
                'PATH_DENIED', 'Linked files or directories are not accepted.')
    real = result.resolve()
    require(real.is_relative_to(base) and real != base, 'PATH_DENIED', 'Path is outside this workspace.')
    if output:
        require(not result.exists(), 'OUTPUT_EXISTS', 'Choose a new output filename; existing files are preserved.')
        require(result.parent.is_dir(), 'OUTPUT_PARENT_MISSING', 'Create the output folder inside the workspace first.')
    else:
        require(result.is_file(), 'INPUT_MISSING', 'Input file does not exist in this workspace.')
        require(result.stat().st_size <= MAX_INPUT, 'INPUT_TOO_LARGE', 'Input exceeds the 64 MiB per-file limit.')
    return result


def check_archive(file):
    if not is_zipfile(file):
        return
    with ZipFile(file) as z:
        infos = z.infolist()
        require(len(infos) <= 12000 and sum(x.file_size for x in infos) <= 256 * 1024 * 1024,
                'ARCHIVE_LIMIT', 'Expanded document is too large.')
        require(all(not x.flag_bits & 1 and x.file_size <= 64 * 1024 * 1024
                    and x.file_size <= max(1024 * 1024, x.compress_size * 1000) for x in infos),
                'ARCHIVE_LIMIT', 'Encrypted or excessively compressed document refused.')
        for item in infos:
            if item.filename.endswith(('.xml', '.rels')):
                content = z.read(item)
                require(b'<!DOCTYPE' not in content.upper() and b'<!ENTITY' not in content.upper(),
                        'UNSAFE_XML', 'External document entities are not accepted.')


def restricted_network(event, args):
    if event in ('socket.connect', 'socket.connect_ex', 'socket.getaddrinfo', 'socket.bind'):
        raise OfficeError('NETWORK_DISABLED', 'Office processing does not access the network.')


def libreoffice():
    candidates = [ROOT / 'libreoffice/program/soffice.exe', ROOT / 'libreoffice/program/soffice',
                  Path('/usr/bin/libreoffice')]
    return next((x for x in candidates if x.is_file()), None)


def convert_with_lo(source, destination):
    binary = libreoffice()
    require(binary, 'CONVERTER_UNAVAILABLE', 'This installation lacks the offline LibreOffice conversion component.')
    allowed = {'.pdf', '.docx', '.xlsx', '.pptx', '.odt', '.ods', '.odp', '.txt', '.csv', '.html', '.rtf', '.doc', '.xls', '.ppt'}
    require(destination.suffix.lower() in allowed, 'FORMAT_UNSUPPORTED', 'Unsupported conversion target.')
    require(source.suffix.lower() in allowed, 'FORMAT_UNSUPPORTED', 'Unsupported conversion source.')
    # Reject remote resources before invoking the converter. This supplements
    # link/macro settings; it is not an OS-level sandbox for LibreOffice.
    if is_zipfile(source):
        with ZipFile(source) as z:
            for name in z.namelist():
                if name.endswith('.rels'):
                    from defusedxml.ElementTree import fromstring
                    for rel in fromstring(z.read(name)):
                        require(rel.get('TargetMode') != 'External' or (rel.get('Type') or '').endswith('/hyperlink'),
                                'EXTERNAL_RESOURCE', 'Convert linked resources to embedded content before conversion.')
                if name.endswith('.xml'):
                    from defusedxml.ElementTree import fromstring
                    content = z.read(name)
                    require(not re.search(rb'(?i)(DDEAUTO|INCLUDETEXT|INCLUDEPICTURE|WEBSERVICE\s*\()', content),
                            'EXTERNAL_RESOURCE', 'Active external-data fields are not accepted.')
                    for element in fromstring(content).iter():
                        for key, value in element.attrib.items():
                            if key.endswith('}href') and element.tag.rsplit('}',1)[-1] not in ('a',):
                                require(not re.match(r'(?i)(?:[a-z][a-z0-9+.-]*:|/|\\\\|\.\./)', value),
                                        'EXTERNAL_RESOURCE', 'External image/data references are not accepted.')
    else:
        raw = source.read_bytes()
        plain = raw.replace(b'\x00', b'')
        require(not re.search(rb'(?i)(DDEAUTO|INCLUDETEXT|INCLUDEPICTURE|WEBSERVICE\s*\(|https?[:\\]|file:|ftp:|\\\\[a-z0-9]|\\objdata)', plain),
                'EXTERNAL_RESOURCE', 'This legacy/HTML/RTF file contains possible external resources; embed or remove them first.')
    with tempfile.TemporaryDirectory(prefix='megmeet-lo-') as folder:
        scratch = Path(folder)
        profile = scratch / 'profile'
        (profile / 'user').mkdir(parents=True)
        (profile / 'user/registrymodifications.xcu').write_text('''<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry"><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item><item oor:path="/org.openoffice.Office.Calc/Content/Update"><prop oor:name="Link" oor:op="fuse"><value>1</value></prop></item><item oor:path="/org.openoffice.Office.Writer/Content/Update"><prop oor:name="Link" oor:op="fuse"><value>2</value></prop></item><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="DisableMacrosExecution" oor:op="fuse"><value>true</value></prop></item></oor:items>''', encoding='utf-8')
        incoming = scratch / ('input' + source.suffix.lower())
        shutil.copyfile(source, incoming)
        outdir = scratch / 'out'
        outdir.mkdir()
        args = [str(binary), '-env:UserInstallation=' + profile.as_uri(), '--headless', '--nologo',
                '--nodefault', '--norestore', '--nofirststartwizard', '--convert-to',
                destination.suffix[1:].lower(), '--outdir', str(outdir), str(incoming)]
        environment = {k: v for k, v in os.environ.items() if k in
                       {'PATH', 'SystemRoot', 'WINDIR', 'TEMP', 'TMP', 'HOME', 'LANG', 'LD_LIBRARY_PATH'}}
        process = subprocess.Popen(args, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                                   env=environment, creationflags=0x08000000 if os.name == 'nt' else 0)
        try:
            process.communicate(timeout=90)
        except subprocess.TimeoutExpired:
            # Windows soffice.exe creates soffice.bin. Kill only our owned tree.
            if os.name == 'nt':
                subprocess.run([str(Path(os.environ.get('SystemRoot', 'C:/Windows'))/'System32/taskkill.exe'),
                                '/PID', str(process.pid), '/T', '/F'], capture_output=True, timeout=10,
                               creationflags=0x08000000)
                process.communicate(timeout=5)
                raise OfficeError('CONVERSION_TIMEOUT', 'Conversion exceeded its time limit; original preserved.')
            # Walk only this still-owned child's descendants, never process names.
            descendants = []
            def children(pid):
                with contextlib.suppress(OSError):
                    for value in Path(f'/proc/{pid}/task/{pid}/children').read_text().split():
                        child = int(value)
                        children(child)
                        descendants.append(child)
            children(process.pid)
            for pid in descendants:
                with contextlib.suppress(ProcessLookupError):
                    os.kill(pid, signal.SIGKILL)
            process.kill()
            process.communicate(timeout=5)
            raise OfficeError('CONVERSION_TIMEOUT', 'Conversion exceeded its time limit; original preserved.')
        created = outdir / ('input' + destination.suffix.lower())
        require(process.returncode == 0 and created.is_file(), 'CONVERSION_FAILED',
                'Office conversion failed; original file has not been changed.')
        require(created.stat().st_size <= MAX_OUTPUT, 'OUTPUT_TOO_LARGE', 'Converted file exceeds the output limit.')
        shutil.copyfile(created, destination)


def ocr_image(image):
    from rapidocr import RapidOCR
    models = ROOT / 'ocr-models'
    config = ROOT / 'ocr-config.json'
    require(config.is_file(), 'OCR_UNAVAILABLE', 'Offline OCR models are not installed.')
    params = json.loads(config.read_text(encoding='utf-8'))
    for name in ('Det.model_path', 'Rec.model_path', 'Cls.model_path'):
        params[name] = str(models / params[name])
        require(Path(params[name]).is_file(), 'OCR_UNAVAILABLE', 'An offline OCR model is missing.')
    params.update({'Global.log_level': 'error', 'EngineConfig.onnxruntime.intra_op_num_threads': 1,
                   'EngineConfig.onnxruntime.inter_op_num_threads': 1})
    engine = RapidOCR(params=params)
    result = engine(image)
    return '\n'.join(result.txts or [])[:MAX_TEXT]


def page_slice(total, options, key='start_page', amount='max_pages'):
    start = integer(options.get(key), 1, 1, max(total, 1)) - 1
    limit = integer(options.get(amount), 20, 1, 60)
    return start, min(total, start + limit)


def read_file(file, options, *, inspect=False):
    suffix = file.suffix.lower()
    require(suffix in READABLE, 'FORMAT_UNSUPPORTED', 'Unsupported document format; call capabilities for supported formats.')
    check_archive(file)
    result = {'format': suffix[1:], 'bytes': file.stat().st_size, 'warnings': []}
    text = ''
    if suffix in LEGACY:
        with tempfile.TemporaryDirectory(prefix='megmeet-convert-') as directory:
            converted = Path(directory) / ('converted.' + LEGACY[suffix])
            convert_with_lo(file, converted)
            result.update(read_file(converted, options, inspect=inspect))
            result['source_format'] = suffix[1:]
            result['warnings'].append('Legacy/ODF conversion may change layout; the original was preserved.')
            return result
    if suffix == '.docx':
        from docx import Document
        document = Document(file)
        result.update(paragraph_count=len(document.paragraphs), table_count=len(document.tables))
        if not inspect:
            # Preserve paragraph/table order, and expose indices for targeted edits.
            parts = []
            from docx.table import Table
            from docx.text.paragraph import Paragraph
            for item in document.iter_inner_content():
                if isinstance(item, Paragraph):
                    parts.append(item.text)
                elif isinstance(item, Table):
                    parts.append('\n'.join('\t'.join(cell.text for cell in row.cells) for row in item.rows))
            for section in document.sections:
                parts += [p.text for p in section.header.paragraphs + section.footer.paragraphs if p.text]
            text = '\n'.join(parts)
            result['paragraphs'] = [{'index': n, 'text': p.text[:4000]} for n, p in enumerate(document.paragraphs[:300])]
    elif suffix in {'.xlsx', '.xls'}:
        if suffix == '.xlsx':
            import openpyxl
            workbook = openpyxl.load_workbook(file, read_only=True, data_only=options.get('cached_values') is True, keep_links=False)
            try:
                result['sheets'] = [{'name': s.title, 'rows': s.max_row, 'columns': s.max_column} for s in workbook]
                if not inspect:
                    sheet = workbook[options.get('sheet', workbook.sheetnames[0])]
                    start = integer(options.get('start_row'), 1, 1, 1048576)
                    count = integer(options.get('max_rows'), 200, 1, 2000)
                    columns = integer(options.get('max_columns'), 40, 1, 200)
                    # Legal streaming-generated workbooks may omit dimensions.
                    # Inspect one extra row rather than scanning the whole file.
                    stream = sheet.iter_rows(min_row=start, max_col=min(sheet.max_column or columns, columns), values_only=True)
                    rows = [list(row) for row in itertools.islice(stream, count+1)]
                    stream.close()
                    result.update(sheet=sheet.title, start_row=start, rows=rows[:count],
                                  more_rows=len(rows) > count, more_columns=None if sheet.max_column is None else sheet.max_column > columns)
                    result['warnings'].append('Formula strings are shown unless cached_values:true. Cached results may be stale; convert to a new XLSX to recalculate via LibreOffice.')
            finally:
                workbook.close()
        else:
            import xlrd
            workbook = xlrd.open_workbook(file, on_demand=True)
            result['sheets'] = [{'name': s.name, 'rows': s.nrows, 'columns': s.ncols} for s in workbook.sheets()]
            if not inspect:
                sheet = workbook.sheet_by_name(options.get('sheet', workbook.sheet_names()[0]))
                start = integer(options.get('start_row'), 1, 1, 65536)-1
                count = integer(options.get('max_rows'), 200, 1, 2000)
                columns = integer(options.get('max_columns'), 40, 1, 200)
                result.update(sheet=sheet.name, start_row=start+1,
                              rows=[sheet.row_values(n)[:columns] for n in range(start, min(sheet.nrows, start+count))],
                              more_rows=sheet.nrows > start+count, more_columns=sheet.ncols > columns)
            workbook.release_resources()
        if 'rows' in result:
            text = '\n'.join('\t'.join('' if c is None else str(c) for c in row) for row in result['rows'])
    elif suffix == '.pptx':
        from pptx import Presentation
        presentation = Presentation(file)
        result['slide_count'] = len(presentation.slides)
        if not inspect:
            first, end = page_slice(len(presentation.slides), options)
            slides = []
            for i in range(first, end):
                slide = presentation.slides[i]
                parts = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        parts.append(shape.text)
                    elif shape.has_table:
                        parts.append('\n'.join('\t'.join(c.text for c in row.cells) for row in shape.table.rows))
                if slide.has_notes_slide:
                    parts.append(slide.notes_slide.notes_text_frame.text)
                slides.append({'slide': i+1, 'text': '\n'.join(parts)})
            result.update(slides=slides, more_pages=end < len(presentation.slides))
            text = '\n\n'.join(f"[Slide {s['slide']}]\n{s['text']}" for s in slides)
    elif suffix == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(file)
        require(not reader.is_encrypted, 'ENCRYPTED_DOCUMENT', 'Unlock this PDF in a trusted application first.')
        result['page_count'] = len(reader.pages)
        if not inspect:
            first, end = page_slice(len(reader.pages), options)
            pages = []
            for n in range(first, end):
                content = reader.pages[n].extract_text() or ''
                method = 'embedded-text'
                if options.get('ocr') is True and len(content.strip()) < 20:
                    import pypdfium2 as pdfium
                    from PIL import Image
                    with pdfium.PdfDocument(str(file)) as pdf:
                        page = pdf[n]
                        require(page.get_width()*page.get_height()*4 <= 24_000_000,
                                'IMAGE_TOO_LARGE', 'PDF page exceeds OCR pixel limit.')
                        bitmap = page.render(scale=2)
                        content = ocr_image(bitmap.to_pil())
                        bitmap.close()
                        page.close()
                    method = 'ocr'
                pages.append({'page': n+1, 'text': content[:MAX_TEXT], 'method': method})
            result.update(pages=pages, more_pages=end < len(reader.pages))
            if options.get('tables') is True:
                import pdfplumber
                with pdfplumber.open(file) as pdf:
                    result['tables'] = [{'page': n+1, 'tables': [t[:200] for t in pdf.pages[n].extract_tables()[:10]]}
                                        for n in range(first, end)]
                result['warnings'].append('PDF table boundaries are heuristic; verify merged cells and scanned tables.')
            text = '\n\n'.join(f"[Page {p['page']}]\n{p['text']}" for p in pages)
            if any(not p['text'].strip() for p in pages):
                result['warnings'].append('Some pages have no embedded text; retry with ocr:true for scanned pages.')
    elif suffix in IMAGE_EXT:
        from PIL import Image, ImageOps
        Image.MAX_IMAGE_PIXELS = 24_000_000
        with Image.open(file) as img:
            require(img.width * img.height <= 24_000_000, 'IMAGE_TOO_LARGE', 'Image exceeds 24 megapixels.')
            result.update(width=img.width, height=img.height, mode=img.mode)
            if not inspect:
                text = ocr_image(ImageOps.exif_transpose(img).convert('RGB'))
                result['warnings'].append('OCR extracts visible text; it does not describe objects or guarantee handwriting/table accuracy.')
    elif suffix == '.rtf':
        from striprtf.striprtf import rtf_to_text
        text = rtf_to_text(file.read_text(encoding='utf-8', errors='replace')) if not inspect else ''
    elif suffix in {'.html', '.htm', '.docx'}:
        from markitdown import MarkItDown
        text = MarkItDown(enable_plugins=False).convert_local(str(file)).text_content if not inspect else ''
    elif not inspect:
        raw = file.read_bytes()
        try:
            text = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            try:
                text = raw.decode('gb18030')
                result['warnings'].append('Text decoded as GB18030.')
            except UnicodeDecodeError:
                raise OfficeError('TEXT_ENCODING', 'Use UTF-8 or GB18030 text encoding.')
    if not inspect:
        offset = integer(options.get('offset'), 0, 0, max(len(text), 0))
        limit = integer(options.get('limit'), 30000, 1, MAX_TEXT)
        result.update(text=text[offset:offset+limit], offset=offset, total_chars=len(text),
                      truncated=offset+limit < len(text))
        # One bounded representation of large content, never a hidden duplicate.
        if len(json.dumps(result, default=str, ensure_ascii=False)) > 400000:
            for key in ('pages', 'slides', 'paragraphs', 'rows', 'tables'):
                result.pop(key, None)
            result['warnings'].append('Detailed records omitted to keep the response bounded; request a smaller page/row range.')
    return result


def paragraph_runs_replace(paragraph, old, new):
    """Replace across run boundaries while preserving unaffected run formatting."""
    count = 0
    while old in paragraph.text:
        previous = paragraph.text
        require(count < 10000, 'EDIT_LIMIT', 'Too many replacements in a single paragraph.')
        start = paragraph.text.index(old)
        end = start + len(old)
        position = 0
        inserted = False
        for run in paragraph.runs:
            value = run.text
            stop = position + len(value)
            if position < end and stop > start:
                left = value[:max(0, start-position)]
                right = value[max(0, end-position):] if end < stop else ''
                run.text = left + (new if not inserted else '') + right
                inserted = True
            position = stop
        if not inserted:
            raise OfficeError('UNSUPPORTED_TEXT_STRUCTURE', 'This text spans unsupported fields/hyperlinks; original is preserved.')
        require(paragraph.text != previous, 'UNSUPPORTED_TEXT_STRUCTURE', 'Unsupported text field; original preserved.')
        count += 1
        # A replacement containing the search term must not be replaced again.
        if old in new:
            break
    return count


def word_write(source, target, options):
    from docx import Document
    from docx.shared import Pt
    document = Document(source) if source else Document()
    changes = 0
    for pair in options.get('replacements', []):
        old, new = string(pair['find']), string(pair['replace'])
        require(old and old not in new, message='Replacement must not contain the search string.')
        paragraphs = list(document.paragraphs)
        def collect(tables):
            for table in tables:
                for row in table.rows:
                    for cell in row.cells:
                        paragraphs.extend(cell.paragraphs)
                        collect(cell.tables)
        collect(document.tables)
        for section in document.sections:
            paragraphs.extend(section.header.paragraphs + section.footer.paragraphs)
        for paragraph in paragraphs:
            changes += paragraph_runs_replace(paragraph, old, new)
    for patch in options.get('paragraph_edits', []):
        index = integer(patch.get('index'), None, 0, len(document.paragraphs)-1)
        require(index is not None)
        paragraph = document.paragraphs[index]
        paragraph.text = string(patch['text'])
        changes += 1
    for item in options.get('paragraphs', []):
        if isinstance(item, str):
            document.add_paragraph(string(item))
        else:
            content = string(item['text'])
            if 'heading' in item:
                document.add_heading(content, integer(item['heading'], 1, 0, 9))
            else:
                p = document.add_paragraph()
                r = p.add_run(content)
                r.bold, r.italic = bool(item.get('bold')), bool(item.get('italic'))
                if 'font_size' in item:
                    r.font.size = Pt(integer(item['font_size'], 11, 6, 96))
        changes += 1
    for matrix in options.get('tables', []):
        require(isinstance(matrix, list) and 0 < len(matrix) <= 1000 and 0 < len(matrix[0]) <= 50)
        table = document.add_table(rows=len(matrix), cols=len(matrix[0]))
        table.style = 'Table Grid'
        for row, values in zip(table.rows, matrix):
            require(len(values) == len(matrix[0]))
            for cell, value in zip(row.cells, values):
                cell.text = str(value)
        changes += 1
    for patch in options.get('table_cells', []):
        table = document.tables[integer(patch['table'], None, 0, len(document.tables)-1)]
        row = integer(patch['row'], None, 0, len(table.rows)-1)
        col = integer(patch['column'], None, 0, len(table.columns)-1)
        table.cell(row, col).text = string(patch['text'])
        changes += 1
    require(changes > 0, 'NO_CHANGES', 'No matching text or content changes; original preserved.')
    document.save(target)
    return {'changes': changes, 'warnings': ['Complex fields, tracked changes and embedded objects require a visual review in Office.']}


def spreadsheet_write(source, target, options):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.workbook.properties import CalcProperties
    book = openpyxl.load_workbook(source, keep_links=False) if source else openpyxl.Workbook()
    for entry in options.get('sheets', []):
        title = string(entry['name'], 31)
        require(title and not re.search(r'[\\/*?:\[\]]', title))
        sheet = book[title] if title in book.sheetnames else book.create_sheet(title)
        rows = entry.get('rows', [])
        require(isinstance(rows, list) and len(rows) <= 10000)
        for n, values in enumerate(rows, 1):
            require(isinstance(values, list) and len(values) <= 200)
            for c, value in enumerate(values, 1):
                require(value is None or isinstance(value, (str, int, float, bool)))
                cell = sheet.cell(n, c, value)
                # Only explicitly tagged cells can become formulas.
                if isinstance(value, str) and value.startswith('='):
                    cell.data_type = 's'
        for patch in entry.get('cells', []):
            address = string(patch['cell'], 20)
            require(re.fullmatch(r'[A-Z]{1,3}[1-9][0-9]{0,6}', address))
            cell = sheet[address]
            require(cell.row <= 1048576 and cell.column <= 16384)
            value = patch.get('value')
            require(value is None or isinstance(value, (str, int, float, bool)))
            if 'formula' in patch:
                value = string(patch['formula'], 8192)
                require(value.startswith('=') and not re.search(r'(?i)(WEBSERVICE|HYPERLINK|RTD|DDE|\[|\]|\|)', value),
                        'EXTERNAL_FORMULA_DENIED', 'External-data formulas are not enabled.')
            cell.value = value
            if isinstance(value, str) and value.startswith('=') and 'formula' not in patch:
                cell.data_type = 's'
            if 'number_format' in patch:
                cell.number_format = string(patch['number_format'], 100)
            if 'bold' in patch:
                cell.font = Font(bold=bool(patch['bold']))
            if 'fill' in patch:
                require(re.fullmatch(r'[0-9A-Fa-f]{6}', patch['fill']))
                cell.fill = PatternFill('solid', fgColor=patch['fill'])
        if 'freeze_panes' in entry:
            require(re.fullmatch(r'[A-Z]{1,3}[1-9][0-9]{0,6}', entry['freeze_panes']))
            sheet.freeze_panes = entry['freeze_panes']
        if entry.get('auto_filter'):
            sheet.auto_filter.ref = sheet.dimensions
        for column, width in entry.get('column_widths', {}).items():
            require(re.fullmatch(r'[A-Z]{1,3}', column))
            sheet.column_dimensions[column].width = integer(width, 15, 1, 100)
        for specification in entry.get('charts', []):
            kind = specification.get('type', 'bar')
            require(kind in ('bar', 'line'))
            chart = BarChart() if kind == 'bar' else LineChart()
            chart.title = string(specification.get('title', ''), 200)
            a = integer(specification.get('min_col'), 2, 1, 200)
            b = integer(specification.get('max_col'), a, a, 200)
            last = integer(specification.get('max_row'), sheet.max_row, 2, 10000)
            chart.add_data(Reference(sheet, min_col=a, max_col=b, min_row=1, max_row=last), titles_from_data=True)
            chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=last))
            anchor = specification.get('anchor', 'G2')
            require(re.fullmatch(r'[A-Z]{1,3}[1-9][0-9]{0,6}', anchor))
            sheet.add_chart(chart, anchor)
    require(options.get('sheets'), 'NO_CHANGES', 'Provide at least one worksheet change.')
    if not source and len(book.worksheets) > 1 and not options.get('keep_empty_sheet', False):
        default = book.worksheets[0]
        if default.title == 'Sheet' and default.max_row == 1 and default['A1'].value is None:
            book.remove(default)
    book.calculation = CalcProperties(fullCalcOnLoad=True)
    book.save(target)
    return {'sheets': book.sheetnames, 'warnings': ['Formulas are stored but not evaluated here. Use convert to recalculate or export PDF. Macros and complex Excel features are not supported.']}


def presentation_write(source, target, options):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation(source) if source else Presentation()
    for pair in options.get('replacements', []):
        old, new = string(pair['find']), string(pair['replace'])
        require(old and old not in new)
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        paragraph_runs_replace(paragraph, old, new)
    for entry in options.get('slides', []):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = string(entry.get('title', ''), 500)
        text = slide.placeholders[1].text_frame
        for i, content in enumerate(entry.get('bullets', [])):
            paragraph = text.paragraphs[0] if i == 0 else text.add_paragraph()
            paragraph.text = string(content, 4000)
        if 'notes' in entry:
            slide.notes_slide.notes_text_frame.text = string(entry['notes'])
    require(options.get('slides') or options.get('replacements'), 'NO_CHANGES', 'Provide slides or text replacements.')
    prs.save(target)
    return {'slides': len(prs.slides), 'warnings': ['Complex animations, embedded objects and precise layout require visual review.']}


def pdf_operation(source, target, options, base):
    from pypdf import PdfReader, PdfWriter
    action = options.get('action', 'extract')
    writer = PdfWriter()
    if action == 'merge':
        inputs = options.get('inputs')
        require(isinstance(inputs, list) and 1 <= len(inputs) <= 20)
        for relative in inputs:
            file = workspace_path(base, relative)
            require(file.suffix.lower() == '.pdf')
            reader = PdfReader(file)
            require(not reader.is_encrypted, 'ENCRYPTED_DOCUMENT', 'Unlock PDFs before merging.')
            require(len(writer.pages)+len(reader.pages) <= 500, 'PAGE_LIMIT', 'At most 500 output pages.')
            writer.append(reader, import_outline=False)
    else:
        require(source and source.suffix.lower() == '.pdf')
        reader = PdfReader(source)
        require(not reader.is_encrypted, 'ENCRYPTED_DOCUMENT', 'Unlock PDFs in a trusted application first.')
        require(len(reader.pages) <= 500, 'PAGE_LIMIT', 'At most 500 pages.')
        pages = options.get('pages', list(range(1, len(reader.pages)+1)))
        require(isinstance(pages, list) and 1 <= len(pages) <= 500)
        for n in pages:
            n = integer(n, None, 1, len(reader.pages))
            writer.add_page(reader.pages[n-1])
        if action == 'rotate':
            angle = options.get('degrees', 90)
            require(angle in (90, 180, 270))
            for page in writer.pages:
                page.rotate(angle)
        elif action == 'watermark':
            from reportlab.pdfgen import canvas
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.cidfonts import UnicodeCIDFont
            text = string(options['text'], 200)
            pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
            for page in writer.pages:
                buffer = io.BytesIO()
                c = canvas.Canvas(buffer, pagesize=(float(page.mediabox.width), float(page.mediabox.height)))
                c.setFont('STSong-Light', 22)
                c.setFillColorRGB(0.55, 0.55, 0.55)
                c.translate(float(page.mediabox.width)/2, float(page.mediabox.height)/2)
                c.rotate(35)
                c.drawCentredString(0, 0, text)
                c.save()
                page.merge_page(PdfReader(buffer).pages[0])
        elif action == 'fill_forms':
            writer = PdfWriter(clone_from=reader)
            fields = options.get('fields', {})
            require(isinstance(fields, dict) and len(fields) <= 200)
            known = reader.get_fields() or {}
            require(set(fields) <= set(known), 'UNKNOWN_FIELD', 'Unknown PDF form field.')
            require(all(isinstance(v, str) and len(v) <= 4000 for v in fields.values()))
            for page in writer.pages:
                writer.update_page_form_field_values(page, fields, auto_regenerate=False)
        elif action != 'extract':
            raise OfficeError('ACTION_UNSUPPORTED', 'PDF actions: merge, extract, rotate, watermark, fill_forms.')
    with target.open('wb') as stream:
        writer.write(stream)
    return {'pages': len(writer.pages), 'warnings': ['PDF text reflow/editing and secure redaction are not supported; use source document editing.']}


def image_operation(source, target, options):
    from PIL import Image, ImageOps
    require(source and source.suffix.lower() in IMAGE_EXT and target.suffix.lower() in IMAGE_EXT)
    Image.MAX_IMAGE_PIXELS = 24_000_000
    with Image.open(source) as original:
        require(original.width * original.height <= 24_000_000, 'IMAGE_TOO_LARGE', 'Image exceeds 24 megapixels.')
        image = ImageOps.exif_transpose(original).copy()
    if 'crop' in options:
        crop = options['crop']
        require(isinstance(crop, list) and len(crop) == 4 and all(type(x) is int for x in crop))
        require(0 <= crop[0] < crop[2] <= image.width and 0 <= crop[1] < crop[3] <= image.height)
        image = image.crop(crop)
    if 'width' in options or 'height' in options:
        width = integer(options.get('width'), image.width, 1, 12000)
        height = integer(options.get('height'), round(image.height * width / image.width), 1, 12000)
        require(width*height <= 24_000_000, 'IMAGE_TOO_LARGE', 'Output image exceeds 24 megapixels.')
        image = image.resize((width, height), Image.Resampling.LANCZOS)
    if 'rotate' in options:
        require(options['rotate'] in (90, 180, 270))
        image = image.rotate(options['rotate'], expand=True)
    if target.suffix.lower() in ('.jpg', '.jpeg'):
        background = Image.new('RGB', image.size, 'white')
        if 'A' in image.getbands():
            background.paste(image, mask=image.getchannel('A'))
            image = background
        else:
            image = image.convert('RGB')
    image.save(target, quality=integer(options.get('quality'), 90, 20, 100))
    return {'width': image.width, 'height': image.height, 'metadata': 'not copied'}


def create_pdf(target, options):
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from xml.sax.saxutils import escape
    pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
    styles = getSampleStyleSheet()
    styles['Normal'].fontName = 'STSong-Light'
    styles['Normal'].leading = 16
    content = options.get('paragraphs', [])
    require(isinstance(content, list) and 0 < len(content) <= 1000)
    story = []
    for value in content:
        story.extend([Paragraph(escape(string(value)).replace('\n', '<br/>'), styles['Normal']), Spacer(1, 8)])
    SimpleDocTemplate(str(target)).build(story)
    return {'paragraphs': len(content)}


def capabilities():
    return {'version': VERSION, 'protocol': PROTOCOL, 'read_formats': sorted(x[1:] for x in READABLE),
            'offline_ocr': (ROOT / 'ocr-config.json').is_file(), 'libreoffice': bool(libreoffice()),
            'operations': {
                'inspect': 'input: relative filename; metadata only',
                'read': 'input; options: offset/limit, sheet/start_row/max_rows/max_columns/cached_values, start_page/max_pages, ocr:true for scans, tables:true for PDF tables',
                'ocr': 'input: image or PDF; options: start_page/max_pages',
                'create': 'output: new .docx/.xlsx/.pptx/.pdf/.txt/.md/.csv/.tsv/.json file; format follows output suffix',
                'edit': 'input + new output of SAME modern format; never overwrites the original',
                'convert': 'input + new output; Office/ODF/legacy formats/PDF/text via LibreOffice; .xlsx to new .xlsx recalculates formulas; .md/.txt extracts text; image conversions supported',
                'pdf': 'options.action: merge(inputs relative paths), extract(pages 1-based), rotate(pages,degrees), watermark(text), fill_forms(fields)',
                'image': 'input + new output; options: crop:[left,top,right,bottom], width,height,rotate:90/180/270,quality:20..100'},
            'write_options': {
                'docx': {'paragraphs': ['text', {'text':'Title','heading':1}], 'replacements':[{'find':'old','replace':'new'}], 'paragraph_edits':[{'index':0,'text':'new'}], 'tables':[[['A','B'],['1','2']]], 'table_cells':[{'table':0,'row':1,'column':0,'text':'updated'}]},
                'xlsx': {'sheets':[{'name':'Sheet1','rows':[['Name','Value'],['A',12]],'cells':[{'cell':'C2','formula':'=SUM(B2:B5)','number_format':'0.00','bold':True,'fill':'FFEECC'}], 'freeze_panes':'A2','auto_filter':True,'column_widths':{'A':25},'charts':[{'type':'bar','title':'Values','min_col':2,'max_col':2,'max_row':5,'anchor':'E2'}]}]},
                'pptx': {'slides':[{'title':'Title','bullets':['One','Two'],'notes':'Speaker notes'}], 'replacements':[{'find':'old','replace':'new'}]},
                'pdf': {'paragraphs':['Hello','Chinese text is supported']},
                'text': {'text':'content'}, 'csv/tsv': {'rows':[['Name','Value'],['A',12]]}},
            'limits': {'input_bytes':MAX_INPUT,'output_bytes':MAX_OUTPUT,'text_chars':MAX_TEXT,'pages_per_read':60},
            'limitations':['No macros, arbitrary Python, URLs, shell commands or external-data execution.',
                'OCR can be inaccurate. Review financial/legal values and scanned tables.',
                'Complex formatting/embedded objects/tracked changes may not round-trip. Output is always a new file.',
                'Legacy DOC/XLS/PPT editing requires conversion to DOCX/XLSX/PPTX.',
                'Image OCR is not general visual reasoning; it does not require a vision model.']}


def execute(request):
    require(isinstance(request, dict) and request.get('protocol') == PROTOCOL)
    operation = request.get('operation')
    if operation == 'capabilities':
        return capabilities()
    require(operation in ('inspect','read','ocr','create','edit','convert','pdf','image'))
    base = Path(string(request.get('workspace'), 2048)).resolve(strict=True)
    require(base.is_dir() and base != Path(base.anchor), 'WORKSPACE_REQUIRED', 'A selected workspace is required.')
    source = workspace_path(base, request['input']) if request.get('input') else None
    options = request.get('options', {})
    require(isinstance(options, dict) and len(json.dumps(options, ensure_ascii=False)) <= 1024*1024)
    if operation in ('inspect','read','ocr'):
        require(source is not None)
        if operation == 'ocr':
            require(source.suffix.lower() in IMAGE_EXT | {'.pdf'})
            options = {**options, 'ocr':True}
        return read_file(source, options, inspect=operation == 'inspect')
    destination = workspace_path(base, request.get('output'), output=True)
    require(destination.suffix.lower() in READABLE, 'FORMAT_UNSUPPORTED', 'Output filename needs a supported extension.')
    if source:
        check_archive(source)
    if operation == 'edit':
        require(source and source.suffix.lower() == destination.suffix.lower())
    with tempfile.TemporaryDirectory(prefix='.megmeet-office-', dir=destination.parent) as directory:
        staged = Path(directory) / ('result' + destination.suffix.lower())
        extension = destination.suffix.lower()
        if operation == 'pdf':
            require(extension == '.pdf')
            result = pdf_operation(source, staged, options, base)
        elif operation == 'image':
            result = image_operation(source, staged, options)
        elif operation == 'convert':
            require(source is not None)
            if source.suffix.lower() in IMAGE_EXT and extension in IMAGE_EXT:
                result = image_operation(source, staged, options)
            elif extension in ('.md','.txt'):
                parsed = read_file(source, {**options, 'limit':MAX_TEXT})
                require(not parsed.get('truncated') and not parsed.get('more_pages') and not parsed.get('more_rows'),
                        'CONVERSION_TRUNCATED', 'Text export would be incomplete; select a smaller explicit range or use Office conversion.')
                staged.write_text(parsed['text'], encoding='utf-8')
                result = {'warnings':parsed['warnings']}
            else:
                convert_with_lo(source, staged)
                result = {'warnings':['Conversion can change formatting. Review the new file before replacing any original.']}
        elif extension == '.docx':
            result = word_write(source, staged, options)
        elif extension == '.xlsx':
            result = spreadsheet_write(source, staged, options)
        elif extension == '.pptx':
            result = presentation_write(source, staged, options)
        elif extension == '.pdf' and operation == 'create':
            result = create_pdf(staged, options)
        elif extension in ('.txt','.md','.json','.csv','.tsv'):
            if extension in ('.csv','.tsv') and 'rows' in options:
                rows = options['rows']
                require(isinstance(rows,list) and len(rows) <= 10000)
                with staged.open('w', newline='', encoding='utf-8-sig') as stream:
                    writer = csv.writer(stream, delimiter='\t' if extension == '.tsv' else ',')
                    for row in rows:
                        require(isinstance(row,list) and len(row) <= 200)
                        writer.writerow([("'"+v) if isinstance(v,str) and v.lstrip().startswith(('=','+','-','@')) else v for v in row])
            else:
                text = string(options.get('text'))
                if extension == '.json':
                    json.loads(text)
                staged.write_text(text, encoding='utf-8')
            result = {}
        else:
            raise OfficeError('WRITE_UNSUPPORTED','Create/edit supports modern Office formats. Convert legacy files first.')
        require(staged.is_file() and 0 < staged.stat().st_size <= MAX_OUTPUT,
                'OUTPUT_INVALID', 'Office output is empty or exceeds the size limit.')
        # Atomic publish with fail-if-exists. A race cannot overwrite a user file.
        workspace_path(base, request['output'], output=True)
        try:
            os.link(staged, destination)
        except FileExistsError:
            raise OfficeError('OUTPUT_EXISTS','Output already exists; original preserved.')
        result.update(output=destination.relative_to(base).as_posix(), bytes=destination.stat().st_size,
                      sha256=hashlib.sha256(destination.read_bytes()).hexdigest(), original_preserved=True)
        return result


def main():
    # UTF-8 is independent of the employee's Windows console code page.
    raw = sys.stdin.buffer.read(2*1024*1024+1)
    try:
        require(len(raw) <= 2*1024*1024, 'REQUEST_TOO_LARGE', 'Office request exceeds 2 MiB.')
        request = json.loads(raw.decode('utf-8'))
        sys.addaudithook(restricted_network)
        with contextlib.redirect_stdout(sys.stderr):
            result = execute(request)
        envelope = {'protocol':PROTOCOL,'ok':True,'result':result}
    except OfficeError as exc:
        envelope = {'protocol':PROTOCOL,'ok':False,'error':{'code':exc.code,'message':exc.message}}
    except Exception as exc:
        # Exception messages can contain private document text or absolute paths.
        envelope = {'protocol':PROTOCOL,'ok':False,'error':{'code':'DOCUMENT_PROCESSING_FAILED',
                    'message':'Could not process this document. Check format, parameters and encryption; original unchanged.',
                    'error_class':type(exc).__name__}}
    try:
        data = json.dumps(envelope, ensure_ascii=False, default=str, allow_nan=False).encode('utf-8')
        require(len(data) <= 2*1024*1024, 'RESPONSE_TOO_LARGE', 'Use smaller page, row or text ranges.')
    except (ValueError, OfficeError):
        data = json.dumps({'protocol': PROTOCOL, 'ok': False, 'error': {
            'code':'RESPONSE_LIMIT', 'message':'Use smaller ranges; unsupported non-finite values cannot be returned.'}}).encode()
    sys.stdout.buffer.write(data+b'\n')


if __name__ == '__main__':
    main()
