"""Local Docker smoke; only disposable containers, synthetic files and DB rows.

Run with Python 3 on a machine with Docker, passing the freshly built image.
No host ports, no company endpoints, no credentials or document mounts.
"""
import argparse
import json
import subprocess
import time
import uuid


def run(*args, input=None, timeout=45):
    result = subprocess.run(args, input=input, text=True, capture_output=True, timeout=timeout)
    if result.returncode:
        raise RuntimeError("LOCAL_DOCKER_STEP_FAILED")
    return result.stdout.strip()


IMPORTS_AND_FILES = r'''
import importlib,io,json,pkgutil
from pathlib import Path
import app.main,maimi_projects
modules=[]
for found in pkgutil.walk_packages(maimi_projects.__path__,maimi_projects.__name__+'.'):
    importlib.import_module(found.name)
    modules.append(found.name)
from openpyxl import Workbook
from maimi_projects.card_importer import parse_workbook
workbook=Workbook()
workbook.active.append(['工号','姓名','英文名'])
workbook.active.append(['M99999999','演示用户','Example'])
stream=io.BytesIO();workbook.save(stream)
rows=parse_workbook(stream.getvalue())
assert len(rows)==1 and rows[0]['employee_no']=='M99999999'
from docx import Document
document=Document();document.add_paragraph('Synthetic roundtrip')
stream=io.BytesIO();document.save(stream)
assert Document(io.BytesIO(stream.getvalue())).paragraphs[0].text=='Synthetic roundtrip'
import qrcode
from PIL import Image
stream=io.BytesIO();qrcode.make('synthetic-local-only').save(stream,format='PNG')
with Image.open(io.BytesIO(stream.getvalue())) as image:
    assert image.width>0 and image.height>0
from fontTools.ttLib import TTFont
assert callable(TTFont)
import pypdfium2 as pdfium
with pdfium.PdfDocument.new() as pdf:
    page=pdf.new_page(128,72)
    page.close()
    stream=io.BytesIO();pdf.save(stream)
with pdfium.PdfDocument(stream.getvalue()) as pdf:
    page=pdf[0];bitmap=page.render(scale=1)
    try:
        image=bitmap.to_pil()
        try: assert image.size==(128,72)
        finally: image.close()
    finally: bitmap.close();page.close()
print(json.dumps({'status':'PASSED','imported_business_modules':len(modules),
 'synthetic_xlsx_parser':True,'synthetic_docx_roundtrip':True,'synthetic_qr_png':True,
 'synthetic_pdf_raster':True,'company_business_calls':0,'company_assets_loaded':False}))
'''


HTTP_CHECK = r'''
import json,urllib.request,urllib.error
result={}
opener=urllib.request.build_opener(urllib.request.ProxyHandler({}))
for path in ('health','ready'):
    try:
        with opener.open('http://127.0.0.1:8080/'+path,timeout=3) as response:
            result[path]={'status':response.status,'body':json.load(response)}
    except urllib.error.HTTPError as error:
        result[path]={'status':error.code,'body':json.load(error)}
print(json.dumps(result))
'''


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--image',required=True)
    args=parser.parse_args()
    prefix='maimi-r65-gateway-smoke-'+uuid.uuid4().hex[:10]
    network=prefix+'-net'
    names=[]
    network_created=False
    phase='IMAGE_DEPENDENCIES_AND_SYNTHETIC_FILES'
    try:
        files=json.loads(run('docker','run','--rm','-i','--network','none','--entrypoint','python',
                             args.image,'-',input=IMPORTS_AND_FILES))
        run('docker','network','create','--internal',network)
        network_created=True
        phase='LOCAL_DATABASE_START'
        db=prefix+'-db';names.append(db)
        run('docker','run','-d','--name',db,'--network',network,'--network-alias','test-db',
            '--memory','512m','--pids-limit','128','-e','POSTGRES_USER=synthetic',
            '-e','POSTGRES_PASSWORD=local-test-only','-e','POSTGRES_DB=synthetic','postgres:16')
        for attempt in range(50):
            result=subprocess.run(['docker','exec',db,'pg_isready','-U','synthetic'],capture_output=True)
            if result.returncode==0: break
            time.sleep(.2)
        else: raise RuntimeError('LOCAL_DATABASE_NOT_READY')
        redis=prefix+'-redis';names.append(redis)
        run('docker','run','-d','--name',redis,'--network',network,'--network-alias','test-redis',
            '--memory','128m','--pids-limit','128','redis:7','redis-server','--save','','--appendonly','no')
        gateway=prefix+'-gateway';names.append(gateway)
        run('docker','run','-d','--name',gateway,'--network',network,'--init',
            '--memory','256m','--pids-limit','128','-e','DSH_ENABLED=false','-e','AUTH_MODE=iam_direct',
            '-e','DATABASE_URL=postgresql://synthetic:local-test-only@test-db:5432/synthetic',
            '-e','REDIS_URL=redis://test-redis:6379/0',args.image)

        def check():
            return json.loads(run('docker','exec','-i',gateway,'python','-',input=HTTP_CHECK))
        def sql(statement):
            run('docker','exec','-i',db,'psql','-U','synthetic','-d','synthetic','-v','ON_ERROR_STOP=1',
                input=statement)
        phase='WAITING_LOCAL_GATEWAY'
        for attempt in range(40):
            try: first=check();break
            except RuntimeError: time.sleep(.25)
        else: raise RuntimeError('LOCAL_GATEWAY_NOT_READY')
        assert first['health']['status']==200 and first['health']['body']['status']=='healthy'
        assert first['ready']['status']==503
        phase='REAL_DICT_ROW_READINESS_REGRESSION'
        sql('CREATE SCHEMA platform; CREATE TABLE platform.schema_migrations(version text);'
            'CREATE TABLE platform.messages(id int);')
        partial=check();assert partial['ready']['status']==503
        sql('CREATE TABLE platform.users(id int); CREATE TABLE platform.conversations(id int);')
        no_migration=check();assert no_migration['ready']['status']==503
        sql("INSERT INTO platform.schema_migrations VALUES ('005_history_write.sql');")
        complete=check();assert complete['ready']['status']==200
        sql('DROP TABLE platform.users;')
        missing=check();assert missing['ready']['status']==503
        print(json.dumps({'schema':'r65-gateway-fresh-image-local-smoke-v1','status':'PASSED',
            'image':args.image,'files':files,'health_http':200,'ready_empty_schema':503,
            'ready_partial_schema':503,'ready_missing_migration':503,'ready_complete_schema':200,
            'ready_removed_users_table':503,'network':'DISPOSABLE_INTERNAL',
            'published_host_ports':0,'company_requests':0,'production_deployment':False}))
    except Exception as error:
        print(json.dumps({'schema':'r65-gateway-fresh-image-local-smoke-v1','status':'FAILED',
                          'phase':phase,'category':type(error).__name__,'company_requests':0}))
        raise SystemExit(1)
    finally:
        for name in reversed(names):
            subprocess.run(['docker','rm','-f',name],capture_output=True,timeout=20)
        if network_created:
            subprocess.run(['docker','network','rm',network],capture_output=True,timeout=20)


if __name__=='__main__':
    main()
