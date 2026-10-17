"""Bounded adapter around the reviewed meeting-document-studio pure modules.

No original FastAPI app, startup seeding, storage directory or business records
are imported. Company templates are read from an operator-pinned read-only root.
"""
from __future__ import annotations

import hashlib
import io
import json
import os
from pathlib import Path, PurePosixPath
import re
import stat
import zipfile
from datetime import date, datetime

from openpyxl import load_workbook

from .meeting_source.board_signatures import parse_board_excel, SignatureGenerationError
from .meeting_source.meeting_documents import generate_complete_meeting_package, MEETING_KIND_ORDER

MAX_INPUT = 25 * 1024 * 1024
MAX_DOCX = 8 * 1024 * 1024
MAX_OUTPUT = 96 * 1024 * 1024
MAX_ARCHIVE = 32 * 1024 * 1024
KINDS = {
    'board': '董事会', 'independent_directors': '独立董事专门会议',
    'strategy_committee': '战略委员会', 'audit_committee': '审计委员会',
    'nomination_committee': '提名委员会', 'compensation_committee': '薪酬与考核委员会',
}
CATEGORIES = {'signature': '签字页', 'notice': '通知', 'ballot': '表决办法及表决票',
              'resolution': '决议', 'record': '会议记录'}


class MeetingFailure(ValueError):
    pass


def check_zip(data, *, docx=False):
    """Check OOXML expansion before a document library is allowed to open it."""
    if not isinstance(data, bytes) or not data or len(data) > (MAX_DOCX if docx else MAX_INPUT):
        raise MeetingFailure('meeting_file_size_invalid')
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            if not 1 <= len(members) <= 2048 or sum(item.file_size for item in members) > 64 * 1024 * 1024:
                raise MeetingFailure('meeting_archive_limit')
            names = set()
            for item in members:
                parts = PurePosixPath(item.filename).parts
                if (item.filename in names or not parts or item.filename.startswith('/') or
                    '\\' in item.filename or any(x in ('.', '..') for x in parts) or item.flag_bits & 1 or
                    item.file_size > 8 * 1024 * 1024):
                    raise MeetingFailure('meeting_archive_invalid')
                names.add(item.filename)
                if item.filename.endswith(('.xml', '.rels')):
                    xml = archive.read(item)
                    if re.search(br'<!\s*(?:DOCTYPE|ENTITY)\b', xml, re.I):
                        raise MeetingFailure('meeting_xml_entities_refused')
                if item.filename.startswith('xl/externalLinks/'):
                    raise MeetingFailure('meeting_external_links_refused')
            required = 'word/document.xml' if docx else 'xl/workbook.xml'
            if required not in names or '[Content_Types].xml' not in names:
                raise MeetingFailure('meeting_archive_invalid')
    except (zipfile.BadZipFile, RuntimeError, OSError) as exc:
        raise MeetingFailure('meeting_archive_invalid') from exc


def inspect_input(data, input_path):
    """Return the real source parser's package, plus visible default warnings."""
    check_zip(data)
    workbook = None
    warnings = []
    try:
        workbook = load_workbook(io.BytesIO(data), data_only=False, read_only=True, keep_links=False)
        if '三会议案清单' not in workbook.sheetnames:
            raise MeetingFailure('meeting_sheet_missing')
        sheet = workbook['三会议案清单']
        if not 5 <= sheet.max_row <= 104 or not 4 <= sheet.max_column <= 32:
            raise MeetingFailure('meeting_sheet_limit')
        rows = list(sheet.iter_rows())
        for row in rows:
            for cell in row:
                if cell.data_type == 'f':
                    raise MeetingFailure('meeting_formula_values_required')
                if isinstance(cell.value, str) and (len(cell.value) > 8000 or '\x00' in cell.value):
                    raise MeetingFailure('meeting_cell_limit')
        input_path.write_bytes(data)
        package = parse_board_excel(input_path)
        if len(package.memo_title) > 240 or not package.memo_title:
            raise MeetingFailure('meeting_memo_title_invalid')
        for meeting in package.meetings.values():
            if (not 1 <= len(meeting.agendas) <= 100 or len(meeting.term) > 12 or len(meeting.sequence) > 12 or
                sum(len(agenda.children) for agenda in meeting.agendas) > 200 or
                not 1 <= int(meeting.meeting_hour) <= 12 or not 0 <= int(meeting.meeting_minute) <= 59):
                raise MeetingFailure('meeting_fields_limit')
            if meeting.notice_date > meeting.meeting_date:
                warnings.append(f'{KINDS[meeting.kind]}通知日期晚于召开日期，请核对原表。')
        # The reviewed project defaults date-only input to 09:00. Show this
        # rule explicitly; do not let the model present it as supplied data.
        for column in range(4, sheet.max_column + 1):
            header, when = rows[3][column - 1].value, rows[2][column - 1].value
            if header and when and ((isinstance(when, str) and not re.search(r'\d\s*[:：点时]', when)) or
                                    (isinstance(when, date) and not isinstance(when, datetime)) or
                                    (isinstance(when, datetime) and when.hour == when.minute == 0) or
                                    isinstance(when, (int, float))):
                warnings.append('原表有未明确时分的日期输入，当前展示时间沿用原项目解析规则；请补充或核对时间后生成。')
                break
        return package, warnings
    except MeetingFailure:
        raise
    except (SignatureGenerationError, ValueError, TypeError, KeyError, IndexError, OverflowError) as exc:
        raise MeetingFailure('meeting_input_invalid') from exc
    finally:
        if workbook is not None:
            workbook.close()


def _regular_bytes(root, relative, maximum):
    # The root is an operator mount, never a model argument. Refuse symlinks
    # in each component and confirm identity again after opening/reading.
    root = Path(root)
    target = root
    for part in ('', *PurePosixPath(relative).parts):
        target = target / part
        info = target.lstat()
        if stat.S_ISLNK(info.st_mode):
            raise MeetingFailure('meeting_template_link_refused')
    before = target.stat()
    if not stat.S_ISREG(before.st_mode) or before.st_size > maximum or before.st_nlink != 1:
        raise MeetingFailure('meeting_template_file_invalid')
    with os.fdopen(os.open(target, os.O_RDONLY | getattr(os, 'O_NOFOLLOW', 0)), 'rb') as handle:
        opened = os.fstat(handle.fileno())
        # CPython/Windows path.stat and fstat may expose birth/change time
        # differently. POSIX keeps the stronger ctime comparison; both paths
        # also require the operator-pinned complete content digest below.
        identity = lambda s: (s.st_dev, s.st_ino, s.st_size, s.st_mtime_ns,
                              s.st_ctime_ns if os.name != 'nt' else None)
        if identity(before) != identity(opened):
            raise MeetingFailure('meeting_template_changed')
        data = handle.read(maximum + 1)
        if len(data) > maximum or identity(opened) != identity(os.fstat(handle.fileno())):
            raise MeetingFailure('meeting_template_changed')
    return data


class MeetingAssets:
    """Exact deployment manifest of 8 signature + 6x4 meeting template paths."""
    def __init__(self, root, manifest):
        self.root = Path(root)
        if not isinstance(manifest, dict) or set(manifest) != {'schema', 'files'} or manifest['schema'] != 'meeting-assets-v1':
            raise MeetingFailure('meeting_template_manifest_invalid')
        entries = manifest['files']
        if not isinstance(entries, list) or len(entries) != 32:
            raise MeetingFailure('meeting_template_manifest_invalid')
        self.files = []
        seen, kinds = set(), {kind: set() for kind in MEETING_KIND_ORDER}
        signatures = 0
        for entry in entries:
            if not isinstance(entry, dict) or set(entry) != {'path', 'bytes', 'sha256'}:
                raise MeetingFailure('meeting_template_manifest_invalid')
            relative = entry['path']
            if (not isinstance(relative, str) or len(relative) > 240 or '\\' in relative or
                any(c in relative for c in '\x00\r\n:') or relative.startswith('/') or
                any(p in ('', '.', '..') for p in relative.split('/')) or relative in seen or
                not relative.endswith('.docx') or not re.fullmatch(r'[0-9a-f]{64}', str(entry['sha256'])) or
                type(entry['bytes']) is not int or not 1 <= entry['bytes'] <= MAX_DOCX):
                raise MeetingFailure('meeting_template_manifest_invalid')
            parts = relative.split('/')
            if len(parts) == 2 and parts[0] == 'board-signatures':
                signatures += 1
            elif len(parts) >= 3 and parts[0] == 'meeting-templates' and parts[1] in kinds:
                prefix = parts[-1][:2]
                if prefix not in ('01', '02', '03', '04') or prefix in kinds[parts[1]]:
                    raise MeetingFailure('meeting_template_manifest_invalid')
                kinds[parts[1]].add(prefix)
            else:
                raise MeetingFailure('meeting_template_manifest_invalid')
            seen.add(relative)
            self.files.append(dict(entry))
        if (signatures != 8 or any(len(value) != 4 for value in kinds.values()) or
                sum(entry['bytes'] for entry in self.files) > 64 * 1024 * 1024 or
                any(len(entry['path'].split('/')[-1]) > 128 for entry in self.files)):
            raise MeetingFailure('meeting_template_manifest_invalid')
        self.revision = hashlib.sha256(json.dumps(self.files, sort_keys=True).encode()).hexdigest()

    def stage(self, destination):
        """Copy only pinned templates into one private task temp directory."""
        for entry in self.files:
            try:
                data = _regular_bytes(self.root, entry['path'], MAX_DOCX)
            except OSError as exc:
                raise MeetingFailure('meeting_template_file_unavailable') from exc
            if len(data) != entry['bytes'] or hashlib.sha256(data).hexdigest() != entry['sha256']:
                raise MeetingFailure('meeting_template_changed')
            check_zip(data, docx=True)
            target = destination.joinpath(*entry['path'].split('/'))
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(data)


def render_package(data, scratch, assets):
    input_path = scratch / 'meeting.xlsx'
    package, warnings = inspect_input(data, input_path)
    templates = scratch / 'templates'
    assets.stage(templates)
    try:
        actual, output_dir, documents = generate_complete_meeting_package(
            input_path, templates / 'board-signatures', templates / 'meeting-templates', scratch / 'output')
        outputs = []
        total = 0
        for item in documents:
            relative = str(item['filename']).replace('\\', '/')
            path = output_dir.joinpath(*relative.split('/'))
            if not path.resolve().is_relative_to(output_dir.resolve()):
                raise MeetingFailure('meeting_output_invalid')
            content = path.read_bytes()
            total += len(content)
            if len(content) > MAX_DOCX or total > MAX_OUTPUT:
                raise MeetingFailure('meeting_output_limit')
            outputs.append({**item, 'filename': relative, 'name': path.name, 'data': content})
        if not 12 <= len(outputs) <= 32:
            raise MeetingFailure('meeting_output_invalid')
        packed = io.BytesIO()
        with zipfile.ZipFile(packed, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
            for item in outputs:
                archive.writestr(item['filename'], item['data'])
        if packed.tell() > MAX_ARCHIVE:
            raise MeetingFailure('meeting_archive_output_too_large')
        return actual, warnings, outputs, packed.getvalue()
    except SignatureGenerationError as exc:
        raise MeetingFailure('meeting_template_layout_invalid') from exc
