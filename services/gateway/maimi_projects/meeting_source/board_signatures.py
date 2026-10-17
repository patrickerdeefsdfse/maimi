from __future__ import annotations

import json
import re
import shutil
import zipfile
from dataclasses import asdict, dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = f"{{{W_NS}}}"
NS = {"w": W_NS}
CN_NUM = "一二三四五六七八九十零〇"
MEETING_NUMBER = CN_NUM + "0-9"


class SignatureGenerationError(ValueError):
    pass


@dataclass(frozen=True)
class AgendaInfo:
    source_index: str
    title: str
    children: tuple[str, ...]
    selection_note: str
    recusal: bool

    def to_dict(self) -> dict[str, Any]:
        data = asdict(self)
        data["children"] = list(self.children)
        return data


@dataclass(frozen=True)
class MeetingInfo:
    kind: str
    label: str
    term: str
    sequence: str
    notice_date: date
    meeting_date: date
    agenda_count: int
    agendas: tuple[AgendaInfo, ...] = ()
    meeting_period: str = "上午"
    meeting_hour: str = "9"
    meeting_minute: str = "00"

    @property
    def meeting_name(self) -> str:
        return f"第{self.term}届董事会{self.label}第{self.sequence}次会议" if self.label else f"第{self.term}届董事会第{self.sequence}次会议"

    def to_dict(self) -> dict[str, Any]:
        data = asdict(self)
        data["meeting_name"] = self.meeting_name
        data["notice_date"] = self.notice_date.isoformat()
        data["meeting_date"] = self.meeting_date.isoformat()
        data["agendas"] = [agenda.to_dict() for agenda in self.agendas]
        data["meeting_time"] = f"{self.meeting_period}{self.meeting_hour}:{self.meeting_minute}"
        return data


@dataclass(frozen=True)
class ExcelMeetingPackage:
    source_filename: str
    memo_title: str
    meetings: dict[str, MeetingInfo]

    def to_dict(self) -> dict[str, Any]:
        return {
            "source_filename": self.source_filename,
            "memo_title": self.memo_title,
            "meetings": {key: value.to_dict() for key, value in self.meetings.items()},
        }


MEETING_TYPES: tuple[tuple[str, str], ...] = (
    ("independent_directors", "独立董事专门会议"),
    ("audit_committee", "审计委员会"),
    ("nomination_committee", "提名委员会"),
    ("compensation_committee", "薪酬与考核委员会"),
    ("strategy_committee", "战略委员会"),
    ("board", ""),
)

# 四个专门委员会：Excel 中未勾选任何议案且未标注会议通知/召开日期时，
# 视为本次不召开——整列跳过，不生成其会议文件，签字页中也删除对应签署块。
OPTIONAL_COMMITTEE_KINDS = frozenset(
    kind for kind, _ in MEETING_TYPES if kind.endswith("_committee")
)


def _date_value(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        converted = from_excel(value)
        return converted.date() if isinstance(converted, datetime) else converted
    if isinstance(value, str):
        match = re.search(r"(20\d{2})\s*[年/.\-]\s*(\d{1,2})\s*[月/.\-]?\s*(\d{1,2})", value)
        if match:
            return date(*(int(piece) for piece in match.groups()))
    raise SignatureGenerationError(f"无法识别日期：{value!r}")


def _date_value_optional(value: Any) -> date | None:
    """日期缺失或无法识别时返回 None，用于判断委员会是否召开。"""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    try:
        return _date_value(value)
    except SignatureGenerationError:
        return None


def _date_parts(value: date) -> tuple[str, str, str]:
    return str(value.year), str(value.month), str(value.day)


def _meeting_clock(value: Any) -> tuple[str, str, str]:
    if isinstance(value, datetime):
        period = "下午" if value.hour >= 12 else "上午"
        hour = value.hour if value.hour <= 12 else value.hour - 12
        return period, str(hour or 12), f"{value.minute:02d}"
    if isinstance(value, str):
        match = re.search(r"(上午|下午)?\s*(\d{1,2})\s*[:：点时]\s*(\d{1,2})?", value)
        if match:
            hour = int(match.group(2))
            period = match.group(1) or ("下午" if hour >= 12 else "上午")
            display_hour = hour - 12 if period == "下午" and hour > 12 else hour
            return period, str(display_hour or 12), f"{int(match.group(3) or 0):02d}"
    return "上午", "9", "00"


def _agenda_from_cell(source_index: Any, value: Any, selection: Any) -> AgendaInfo:
    raw = str(value or "").strip()
    if not raw:
        raise SignatureGenerationError(f"议案序号 {source_index!r} 已勾选，但议案名称为空。")
    lines = [line.strip() for line in re.split(r"[\r\n]+", raw) if line.strip()]
    title = re.sub(r"\s*【\s*\d+\s*个子议案\s*】\s*$", "", lines[0]).strip()
    children: list[str] = []
    for line in lines[1:]:
        child = re.sub(r"^\s*\d+\s*[.．、]\s*", "", line).strip()
        child = child.rstrip("；;。").strip()
        if child.startswith("《") and child.endswith("》"):
            child = child[1:-1].strip()
        if child:
            children.append(child)

    note = str(selection or "").strip()
    subset = re.search(r"仅审议\s*(\d+)\s*[-—至]\s*(\d+)\s*项?子议案", note)
    if subset and children:
        start, end = int(subset.group(1)), int(subset.group(2))
        children = children[max(0, start - 1):end]
    return AgendaInfo(
        source_index=str(source_index or "").strip(),
        title=title,
        children=tuple(children),
        selection_note=note,
        recusal="回避表决" in note,
    )


def _meeting_kind(header: str) -> tuple[str, str] | None:
    compact = re.sub(r"\s+", "", header)
    if "股东会" in compact:
        return None
    for kind, label in MEETING_TYPES:
        if label and label in compact:
            return kind, label
    return ("board", "") if "董事会" in compact else None


def _cell_selected(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value == 1
    return str(value or "").strip().upper() in {"√", "✓", "Y", "是", "TRUE", "1"} or str(value or "").strip().startswith("√")


def parse_board_excel(excel_path: Path) -> ExcelMeetingPackage:
    try:
        workbook = load_workbook(excel_path, data_only=True, read_only=True)
    except Exception as exc:  # pragma: no cover - library error detail is surfaced to the user
        raise SignatureGenerationError(f"无法打开 Excel：{exc}") from exc
    try:
        if "三会议案清单" not in workbook.sheetnames:
            raise SignatureGenerationError("未找到“三会议案清单”工作表；该系统仅接受已保存的董事会清单格式。")
        sheet = workbook["三会议案清单"]
        memo_raw = str(sheet.cell(1, 1).value or "").strip()
        memo_title = re.sub(r"^重大事项进程备忘录[-—：:]?", "", memo_raw).strip()
        if not memo_title:
            memo_title = memo_raw

        meetings: dict[str, MeetingInfo] = {}
        for column in range(4, sheet.max_column + 1):
            header_value = sheet.cell(4, column).value
            if not header_value:
                continue
            header = str(header_value)
            kind_info = _meeting_kind(header)
            if not kind_info:
                continue
            kind, label = kind_info
            if kind in meetings:
                raise SignatureGenerationError(f"Excel 中存在多个“{kind}”会议列，需先在规则中配置唯一映射。")
            compact = re.sub(r"\s+", "", header)
            agendas = tuple(
                _agenda_from_cell(sheet.cell(row, 1).value, sheet.cell(row, 2).value, sheet.cell(row, column).value)
                for row in range(5, sheet.max_row + 1)
                if _cell_selected(sheet.cell(row, column).value)
            )
            # 专门委员会本次不召开：未勾选任何议案，也未标注通知/召开日期，整列跳过。
            if kind in OPTIONAL_COMMITTEE_KINDS and not agendas and (
                _date_value_optional(sheet.cell(2, column).value) is None
                and _date_value_optional(sheet.cell(3, column).value) is None
            ):
                continue
            term_match = re.search(rf"第([{MEETING_NUMBER}]+)届董事会", compact)
            # 独立董事专门会议的表头为“董事会第五次独立董事专门会议”，
            # 而各专门委员会通常是“委员会第二次会议”。两种写法都只会有一处“第X次”。
            sequence_match = re.search(rf"第([{MEETING_NUMBER}]+)次", compact)
            if not term_match or not sequence_match:
                raise SignatureGenerationError(f"无法从会议列识别届次：{header}")
            notice_date = _date_value(sheet.cell(2, column).value)
            meeting_raw = sheet.cell(3, column).value
            meeting_date = _date_value(meeting_raw)
            period, hour, minute = _meeting_clock(meeting_raw)
            meetings[kind] = MeetingInfo(
                kind,
                label,
                term_match.group(1),
                sequence_match.group(1),
                notice_date,
                meeting_date,
                len(agendas),
                agendas,
                period,
                hour,
                minute,
            )

        if "board" not in meetings:
            raise SignatureGenerationError("Excel 中未识别到董事会会议列。")
        return ExcelMeetingPackage(excel_path.name, memo_title, meetings)
    finally:
        workbook.close()


def seed_board_library(
    source_templates: Path,
    source_rule: Path,
    library_root: Path,
    rules_root: Path,
    source_meeting_templates: Path | None = None,
    meeting_library_root: Path | None = None,
) -> list[str]:
    if not source_templates.exists() or not source_rule.exists():
        raise SignatureGenerationError("初始化来源中缺少签字页母版或董事会处理规则。")
    library_root.mkdir(parents=True, exist_ok=True)
    rules_root.mkdir(parents=True, exist_ok=True)
    copied: list[str] = []
    signature_templates = sorted(source_templates.glob("*.docx"))
    source_names = {template.name for template in signature_templates}
    # Keep an upgraded local library aligned with the packaged seed set. This
    # also removes templates renamed to correct a signer's name.
    for existing in library_root.glob("*.docx"):
        if existing.name not in source_names:
            existing.unlink()
    for template in signature_templates:
        target = library_root / template.name
        shutil.copy2(template, target)
        copied.append(target.name)
    if source_meeting_templates is not None and meeting_library_root is not None:
        if not source_meeting_templates.exists():
            raise SignatureGenerationError("初始化来源中缺少会议通知、表决票、决议或会议记录母版。")
        meeting_library_root.mkdir(parents=True, exist_ok=True)
        for template in sorted(source_meeting_templates.rglob("*.docx")):
            relative = template.relative_to(source_meeting_templates)
            target = meeting_library_root / relative
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(template, target)
            copied.append(str(Path("meeting-templates") / relative))
    shutil.copy2(source_rule, rules_root / "董事会处理规则.docx")
    (rules_root / "board-signature-rule.json").write_text(
        json.dumps(
            {
                "sheet": "三会议案清单",
                "notice_date_row": 2,
                "meeting_date_row": 3,
                "meeting_header_row": 4,
                "behavior": "按会议列生成签字页、通知、表决办法及表决票、决议、会议记录；会议时间取召开日，通知回执及召集人落款取通知日；通知和表决办法正文的子议案只写母议案，且不得保留空的子议案编号；表决票表格再展开子议案，独董会表决票使用满幅表格；回避表决合并后三列、仅保留一个段落并水平垂直居中；会议记录使用4.1、4.2.1点分编号且议程块之间不留空行；未勾选议案且未标注通知/召开日期的专门委员会视为本次不召开：跳过其会议文件并删除签字页中对应签署块；无来源字段沿用母版或留空，议案介绍不得编造。",
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    return copied


def _part_paragraphs(root: ET.Element) -> list[ET.Element]:
    return list(root.findall(".//w:p", NS))


def _paragraph_text(paragraph: ET.Element) -> str:
    return "".join(node.text or "" for node in paragraph.findall(".//w:t", NS))


SIGNATURE_BLOCK_KINDS: tuple[tuple[str, str], ...] = (
    ("独立董事专门会议", "independent_directors"),
    ("审计委员会", "audit_committee"),
    ("提名委员会", "nomination_committee"),
    ("薪酬与考核委员会", "compensation_committee"),
    ("战略委员会", "strategy_committee"),
)


def _meeting_kind_for_text(text: str) -> str | None:
    """按优先级识别段落所属的会议块，与该会议是否召开无关。"""
    for phrase, kind in SIGNATURE_BLOCK_KINDS:
        if phrase in text:
            return kind
    if "重大事项进程备忘录" in text:
        return "memo"
    return "board" if "董事会" in text else None


def _meeting_for_text(text: str, meetings: dict[str, MeetingInfo]) -> MeetingInfo | None:
    kind = _meeting_kind_for_text(text)
    if kind is None or kind == "memo":
        return None
    return meetings.get(kind)


def _placeholder_values(
    text: str,
    meeting: MeetingInfo | None,
    memo_title: str,
    standalone_notice_date: bool = False,
) -> dict[int, str]:
    values: dict[int, str] = {}
    if "重大事项进程备忘录-【】" in text:
        values[text.index("【】", text.index("重大事项进程备忘录-"))] = memo_title
    if meeting is None:
        return values

    for match in re.finditer(r"第【】届", text):
        values[match.start() + 1] = meeting.term
    for match in re.finditer(r"第【】次", text):
        values[match.start() + 1] = meeting.sequence
    for match in re.finditer(rf"第([{CN_NUM}]+)【】次", text):
        values[match.start() + len("第") + len(match.group(1))] = ""
    date_matches = list(re.finditer(r"【】年【】月【】日", text))
    for position, match in enumerate(date_matches):
        next_start = date_matches[position + 1].start() if position + 1 < len(date_matches) else len(text)
        between = text[match.end():next_start]
        source_date = (
            meeting.notice_date
            if "收到" in between or (standalone_notice_date and len(date_matches) == 1)
            else meeting.meeting_date
        )
        cursor = match.start()
        for part in _date_parts(source_date):
            placeholder = text.find("【】", cursor, match.end())
            values[placeholder] = part
            cursor = placeholder + 2
    return values


def _replace_paragraph(
    paragraph: ET.Element,
    meetings: dict[str, MeetingInfo],
    memo_title: str,
    fallback_meeting: MeetingInfo | None = None,
    standalone_notice_date: bool = False,
) -> None:
    original = _paragraph_text(paragraph)
    if "【】" not in original:
        return
    identified_meeting = _meeting_for_text(original, meetings)
    # “会议时间”和召集人签署日期常被单独放在一行，行内没有会议名称。
    # 前者取会议召开日；后者按最新规则取会议通知日。
    standalone_date = bool(re.fullmatch(r"\s*【】年【】月【】日\s*", original))
    meeting = identified_meeting or (
        fallback_meeting if "会议时间" in original or standalone_date else None
    )
    value_map = _placeholder_values(
        original,
        meeting,
        memo_title,
        standalone_notice_date=standalone_notice_date and standalone_date,
    )
    if not value_map:
        return
    # Word often splits one visual placeholder across adjacent w:t nodes.  Work
    # with paragraph-wide character offsets, but write back to the original
    # nodes so the template's runs, fonts, and alignment remain intact.
    removed_positions = {position + 1 for position in value_map}
    offset = 0
    for node in paragraph.findall(".//w:t", NS):
        value = node.text or ""
        rebuilt: list[str] = []
        for local_position, character in enumerate(value):
            position = offset + local_position
            if position in value_map:
                rebuilt.append(value_map[position])
            elif position not in removed_positions:
                rebuilt.append(character)
        node.text = "".join(rebuilt)
        offset += len(value)


def _remove_absent_committee_blocks(root: ET.Element, meetings: dict[str, MeetingInfo]) -> None:
    """从签字页中删除本次不召开的专门委员会的签署块（含其中的分节符段落）。

    签字页母版按会议块组织：同一会议的回执、签到表、表决票/决议/记录签署页
    连续排列。“会议通知送达回执”标题开启下一个会议块，随该块整体删除或
    保留，避免块删除后残留孤立的回执标题；分节符之后新节起始处尚未出现会
    议名的段落向前归属到下一个可识别的会议块。“重大事项进程备忘录”页固定
    保留；委员会段落中的“董事会”字样不视为切回董事会块。
    """
    absent = {kind for kind in OPTIONAL_COMMITTEE_KINDS if kind not in meetings}
    if not absent:
        return
    body = root.find(f"{W}body")
    if body is None:
        return

    elements = [element for element in body if element.tag in (f"{W}p", f"{W}tbl")]

    def element_text(element: ET.Element) -> str:
        return "".join(node.text or "" for node in element.iter(f"{W}t"))

    def next_identified_kind(index: int) -> str | None:
        for cursor in range(index + 1, len(elements)):
            kind = _meeting_kind_for_text(element_text(elements[cursor]))
            if kind is not None:
                return kind
        return None

    attributed: list[str | None] = []
    current: str | None = None
    # 文档开头视同新节起始：封面等尚未出现会议名的段落向前归属。
    section_started = True
    for index, element in enumerate(elements):
        text = element_text(element)
        kind = _meeting_kind_for_text(text)
        if kind == "memo":
            current = "memo"
            attributed.append("memo")
        elif kind is not None:
            if kind != "board" or current is None:
                current = kind
            attributed.append(current)
        elif section_started:
            attributed.append(next_identified_kind(index))
        elif text.strip() == "会议通知送达回执":
            # 回执标题开启下一个会议块，随该块整体删除或保留。
            current = next_identified_kind(index)
            attributed.append(current)
        else:
            attributed.append(current)
        # 包含 sectPr 的段落结束当前节，其后段落属于新节起始。
        if element.tag == f"{W}p" and element.find(f"{W}pPr/{W}sectPr") is not None:
            section_started = True
        elif kind is not None or text.strip():
            section_started = False

    for element, kind in zip(elements, attributed):
        if kind in absent:
            body.remove(element)


def _rewrite_docx(template_path: Path, output_path: Path, package: ExcelMeetingPackage) -> int:
    parts_changed = 0
    with zipfile.ZipFile(template_path, "r") as source, zipfile.ZipFile(output_path, "w") as destination:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if re.fullmatch(r"word/(document|header\d+|footer\d+)\.xml", entry.filename):
                root = ET.fromstring(data)
                _remove_absent_committee_blocks(root, package.meetings)
                before = "".join(node.text or "" for node in root.iter(W + "t"))
                current_meeting: MeetingInfo | None = None
                recent_texts: list[str] = []
                for paragraph in _part_paragraphs(root):
                    paragraph_text = _paragraph_text(paragraph)
                    identified = _meeting_for_text(paragraph_text, package.meetings)
                    if identified is not None:
                        current_meeting = identified
                    _replace_paragraph(
                        paragraph,
                        package.meetings,
                        package.memo_title,
                        current_meeting,
                        standalone_notice_date="召集人" in "".join(recent_texts[-8:]),
                    )
                    if paragraph_text.strip():
                        recent_texts.append(paragraph_text)
                after = "".join(node.text or "" for node in root.iter(W + "t"))
                if before != after:
                    parts_changed += 1
                    data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            destination.writestr(entry, data)
    return parts_changed


def _placeholder_count(docx_path: Path) -> int:
    """Count unresolved placeholders that should have come from the Excel.

    Other 【】 markers are deliberate user-fill prompts for data that has no
    source. They are allowed to remain in the generated document.
    """
    count = 0
    with zipfile.ZipFile(docx_path) as package:
        for name in package.namelist():
            if re.fullmatch(r"word/(document|header\d+|footer\d+)\.xml", name):
                root = ET.fromstring(package.read(name))
                for paragraph in _part_paragraphs(root):
                    text = _paragraph_text(paragraph)
                    if "【】" not in text:
                        continue
                    required = (
                        "第【】届" in text
                        or "第【】次" in text
                        or "重大事项进程备忘录-【】" in text
                        or ("会议时间" in text and "【】年【】月【】日" in text)
                        or bool(re.fullmatch(r"\s*【】年【】月【】日\s*", text))
                        or ("收到" in text and "【】年【】月【】日" in text)
                        or ("出席会议" in text and "【】年【】月【】日" in text)
                    )
                    if required:
                        count += text.count("【】")
    return count


def generate_signature_pages(excel_path: Path, template_root: Path, output_root: Path) -> tuple[ExcelMeetingPackage, Path, list[dict[str, Any]]]:
    package = parse_board_excel(excel_path)
    templates = sorted(template_root.glob("*.docx"))
    if not templates:
        raise SignatureGenerationError("签字页模板库为空，请先由管理员导入母版。")
    base_job_name = f"{datetime.now():%Y%m%d-%H%M%S}-{package.meetings['board'].meeting_name}"
    for attempt in range(1, 1000):
        job_name = base_job_name if attempt == 1 else f"{base_job_name}-{attempt}"
        job_dir = output_root / job_name
        try:
            job_dir.mkdir(parents=True, exist_ok=False)
            break
        except FileExistsError:
            continue
    else:  # pragma: no cover - defensive limit for an abnormal output directory
        raise SignatureGenerationError("同一会议的生成任务过多，请稍后重试。")

    documents: list[dict[str, Any]] = []
    zip_path = output_root / f"{job_name}-签字页文件包.zip"
    try:
        for template in templates:
            output_name = f"{package.meetings['board'].meeting_name}-{template.stem}.docx"
            output_path = job_dir / output_name
            changed_parts = _rewrite_docx(template, output_path, package)
            if changed_parts == 0:
                raise SignatureGenerationError(f"{template.name} 未发现可替换的会议字段，已停止交付。")
            remaining = _placeholder_count(output_path)
            if remaining:
                raise SignatureGenerationError(f"{template.name} 仍有 {remaining} 个【】占位符，已停止交付。")
            documents.append({"template": template.name, "filename": output_name, "changed_parts": changed_parts})
        with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for item in documents:
                file_path = job_dir / item["filename"]
                archive.write(file_path, arcname=file_path.name)
        return package, zip_path, documents
    except Exception:
        shutil.rmtree(job_dir, ignore_errors=True)
        zip_path.unlink(missing_ok=True)
        raise
