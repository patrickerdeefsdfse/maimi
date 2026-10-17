from __future__ import annotations

import re
import shutil
from copy import deepcopy
from datetime import datetime
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import _Cell, Table
from docx.text.paragraph import Paragraph

from .board_signatures import (
    CN_NUM,
    MEETING_NUMBER,
    AgendaInfo,
    ExcelMeetingPackage,
    MeetingInfo,
    SignatureGenerationError,
    _placeholder_count,
    _rewrite_docx,
    parse_board_excel,
)


MEETING_FOLDER_LABELS = {
    "board": "01-董事会",
    "independent_directors": "02-独立董事专门会议",
    "strategy_committee": "03-战略委员会",
    "audit_committee": "04-审计委员会",
    "nomination_committee": "05-提名委员会",
    "compensation_committee": "06-薪酬与考核委员会",
}
MEETING_KIND_ORDER = tuple(MEETING_FOLDER_LABELS)
WEEKDAYS = "一二三四五六日"
INDEPENDENT_BALLOT_COLUMN_WIDTHS = (6288, 698, 698, 776)
BALLOT_TABLE_INDENT = 108


def _set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    runs = paragraph.runs
    if runs:
        runs[0].text = text
        for run in runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def _strip_numbering(paragraph_element) -> None:
    """Remove Word automatic numbering before writing a manual visible number."""
    properties = paragraph_element.find(qn("w:pPr"))
    if properties is None:
        return
    numbering = properties.find(qn("w:numPr"))
    if numbering is not None:
        properties.remove(numbering)


def _has_direct_numbering(paragraph: Paragraph) -> bool:
    properties = paragraph._p.find(qn("w:pPr"))
    return properties is not None and properties.find(qn("w:numPr")) is not None


def _has_manual_number(text: str) -> bool:
    return bool(re.match(r"\s*(?:[一二三四五六七八九十]+、|\d+(?:\.\d+)*[.、]?\s*)", text))


def _iter_cell_paragraphs(cell: _Cell) -> Iterable[Paragraph]:
    yield from cell.paragraphs
    for table in cell.tables:
        for row in table.rows:
            for nested_cell in row.cells:
                yield from _iter_cell_paragraphs(nested_cell)


def _all_paragraphs(document: Document) -> Iterable[Paragraph]:
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from _iter_cell_paragraphs(cell)
    for section in document.sections:
        for part in (section.header, section.footer):
            yield from part.paragraphs
            for table in part.tables:
                for row in table.rows:
                    for cell in row.cells:
                        yield from _iter_cell_paragraphs(cell)


def _clone_after(source: Paragraph, text: str) -> Paragraph:
    clone = deepcopy(source._p)
    source._p.addnext(clone)
    paragraph = Paragraph(clone, source._parent)
    _set_paragraph_text(paragraph, text)
    return paragraph


def _clone_before(boundary: Paragraph, source_xml, text: str) -> Paragraph:
    clone = deepcopy(source_xml)
    boundary._p.addprevious(clone)
    paragraph = Paragraph(clone, boundary._parent)
    _set_paragraph_text(paragraph, text)
    return paragraph


def _remove_paragraph(paragraph: Paragraph) -> None:
    parent = paragraph._p.getparent()
    if parent is not None:
        parent.remove(paragraph._p)


def _weekday(meeting: MeetingInfo) -> str:
    return WEEKDAYS[meeting.meeting_date.weekday()]


def _cn_index(value: int) -> str:
    if value <= 10:
        return "十" if value == 10 else CN_NUM[value - 1]
    tens, units = divmod(value, 10)
    return (CN_NUM[tens - 1] if tens > 1 else "") + "十" + (CN_NUM[units - 1] if units else "")


def _format_date(meeting: MeetingInfo, notice: bool = False) -> str:
    value = meeting.notice_date if notice else meeting.meeting_date
    return f"{value.year}年{value.month}月{value.day}日"


def _date_for_context(text: str, start: int, meeting: MeetingInfo, document_kind: str) -> str:
    nearby = text[max(0, start - 24):start + 40]
    stripped = text.strip()
    pure_date = bool(re.fullmatch(r"【】年【】月【】日", stripped))
    notice_context = (
        "收到" in nearby
        or "发出通知" in nearby
        or "会议通知于" in nearby
        or "送达" in nearby
        or (document_kind == "notice" and pure_date)
    )
    return _format_date(meeting, notice=notice_context)


def _replace_date_placeholders(text: str, meeting: MeetingInfo, document_kind: str) -> str:
    pattern = re.compile(r"【】年【】月【】日")
    pieces: list[str] = []
    cursor = 0
    for match in pattern.finditer(text):
        pieces.append(text[cursor:match.start()])
        pieces.append(_date_for_context(text, match.start(), meeting, document_kind))
        cursor = match.end()
    pieces.append(text[cursor:])
    return "".join(pieces)


def _replace_hardcoded_dates(text: str, meeting: MeetingInfo, document_kind: str) -> str:
    pattern = re.compile(r"20\d{2}年\d{1,2}月\d{1,2}日")
    return pattern.sub(
        lambda match: _date_for_context(text, match.start(), meeting, document_kind),
        text,
    )


def _replace_meeting_identity(text: str, meeting: MeetingInfo) -> str:
    number = rf"[{MEETING_NUMBER}]+"
    label = re.escape(meeting.label)
    if meeting.kind == "independent_directors":
        patterns = (
            rf"第{number}届董事会第{number}次独立董事专门会议",
            rf"第{number}届董事会独立董事专门会议第{number}次会议",
        )
    elif meeting.label:
        patterns = (rf"第{number}届董事会{label}第{number}次会议",)
    else:
        patterns = (rf"第{number}届董事会第{number}次会议",)
    for pattern in patterns:
        text = re.sub(pattern, meeting.meeting_name, text)
    text = re.sub(r"第【】届", f"第{meeting.term}届", text)
    text = re.sub(r"第【】次", f"第{meeting.sequence}次", text)
    return text


def _replace_metadata(text: str, meeting: MeetingInfo, document_kind: str, memo_title: str) -> str:
    text = _replace_meeting_identity(text, meeting)
    text = _replace_date_placeholders(text, meeting, document_kind)
    text = _replace_hardcoded_dates(text, meeting, document_kind)
    clock = f"{meeting.meeting_period}{meeting.meeting_hour}:{meeting.meeting_minute}"
    text = re.sub(r"(?:上午|下午)?【】\s*[:：]\s*【】", clock, text)
    text = re.sub(r"(?:上午|下午)?【】\s*点", clock, text)
    text = re.sub(r"(上午|下午)\s*\d{1,2}\s*[:：]\s*\d{2}", clock, text)
    text = re.sub(r"星期【】", f"星期{_weekday(meeting)}", text)
    text = re.sub(r"星期[一二三四五六日天]", f"星期{_weekday(meeting)}", text)
    text = text.replace("重大事项进程备忘录-【】", f"重大事项进程备忘录-{memo_title}")
    # Non-empty bracketed text in the meeting templates is authoring guidance
    # (for example 【子议案1】), not a user input field. Remove that guidance,
    # while preserving the exact empty marker 【】 for genuinely source-less
    # fields that the user must complete.
    return re.sub(r"【[^】]+】", "", text)


def _agenda_lines(agendas: tuple[AgendaInfo, ...], verb: str) -> list[str]:
    lines: list[str] = []
    for agenda in agendas:
        if agenda.children:
            lines.append(f"逐项审议通过《{agenda.title}》：")
        else:
            lines.append(f"{verb}《{agenda.title}》")
    for index, line in enumerate(lines):
        if line.endswith("："):
            continue
        lines[index] = line + ("。" if index == len(lines) - 1 else "；")
    return lines


def _replace_agenda_line_placeholders(document: Document, meeting: MeetingInfo) -> None:
    placeholders = [
        paragraph
        for paragraph in document.paragraphs
        if "审议" in paragraph.text
        and "《" in paragraph.text
        and ("【" in paragraph.text or "《》" in paragraph.text)
    ]
    if not placeholders:
        raise SignatureGenerationError("母版中未找到议案清单占位段落。")
    verb = "审议通过" if "审议通过" in placeholders[0].text else "审议"
    lines = _agenda_lines(meeting.agendas, verb)
    anchor = placeholders[-1]
    for index, line in enumerate(lines):
        if index < len(placeholders):
            _set_paragraph_text(placeholders[index], line)
            if _has_manual_number(line):
                _strip_numbering(placeholders[index]._p)
            anchor = placeholders[index]
        else:
            anchor = _clone_after(anchor, line)
            if _has_manual_number(line):
                _strip_numbering(anchor._p)
    for paragraph in placeholders[len(lines):]:
        _remove_paragraph(paragraph)


def _remove_agenda_guidance_paragraphs(document: Document) -> None:
    """Remove numbered child-agenda examples that belong only to the template."""
    for paragraph in list(document.paragraphs):
        compact = re.sub(r"\s+", "", paragraph.text)
        if re.fullmatch(r"【子议案\d*(?:名称)?】[；。]?", compact):
            _remove_paragraph(paragraph)


def _find_boundary(paragraphs: list[Paragraph], start: int, phrases: tuple[str, ...]) -> int:
    for index in range(start + 1, len(paragraphs)):
        if any(phrase in paragraphs[index].text for phrase in phrases):
            return index
    raise SignatureGenerationError(f"母版议案区缺少结束标记：{' / '.join(phrases)}")


def _resolution_vote_text(source: str, agenda: AgendaInfo) -> str:
    return "表决结果：回避表决。" if agenda.recusal else source


def _replace_resolution_blocks(document: Document, meeting: MeetingInfo) -> None:
    paragraphs = list(document.paragraphs)
    start = next(
        (
            index
            for index, paragraph in enumerate(paragraphs)
            if "【" in paragraph.text and "《" in paragraph.text and "审议" in paragraph.text
        ),
        None,
    )
    if start is None:
        raise SignatureGenerationError("决议母版中未找到议案段落。")
    end = _find_boundary(paragraphs, start, ("本页以下无正文", "以下无正文"))
    zone = paragraphs[start:end]
    title_source = deepcopy(zone[0]._p)
    _strip_numbering(title_source)
    intro_paragraph = next((p for p in zone if "议案内容介绍" in p.text), zone[0])
    intro_source = deepcopy(intro_paragraph._p)
    vote_paragraph = next((p for p in zone if "表决结果" in p.text), None)
    if vote_paragraph is None:
        raise SignatureGenerationError("决议母版中未找到表决结果样式。")
    vote_source = deepcopy(vote_paragraph._p)
    vote_text = vote_paragraph.text
    child_paragraph = next((p for p in zone if "子议案" in p.text and "名称" in p.text), zone[0])
    child_source = deepcopy(child_paragraph._p)
    _strip_numbering(child_source)
    boundary = paragraphs[end]
    for paragraph in zone:
        _remove_paragraph(paragraph)

    for index, agenda in enumerate(meeting.agendas, 1):
        action = "逐项审议通过" if agenda.children else "审议通过"
        _clone_before(boundary, title_source, f"{_cn_index(index)}、{action}《{agenda.title}》")
        if agenda.children:
            _clone_before(boundary, title_source, "子议案如下：")
            for child_index, child in enumerate(agenda.children, 1):
                _clone_before(boundary, child_source, f"{child_index}. {child}")
                _clone_before(boundary, intro_source, "")
                _clone_before(boundary, vote_source, _resolution_vote_text(vote_text, agenda))
        else:
            _clone_before(boundary, intro_source, "")
            _clone_before(boundary, vote_source, _resolution_vote_text(vote_text, agenda))


def _replace_record_blocks(document: Document, meeting: MeetingInfo) -> None:
    paragraphs = list(document.paragraphs)
    start = next(
        (
            index
            for index, paragraph in enumerate(paragraphs)
            if "【" in paragraph.text and "《" in paragraph.text and "审议" in paragraph.text
        ),
        None,
    )
    if start is None:
        raise SignatureGenerationError("会议记录母版中未找到议案段落。")
    end = _find_boundary(paragraphs, start, ("主持人宣读会议决议", "主持人宣布会议结束"))
    zone = paragraphs[start:end]
    title_source = deepcopy(zone[0]._p)
    _strip_numbering(title_source)
    result_paragraph = next((p for p in zone if "投票表决" in p.text or "审议通过了" in p.text), None)
    if result_paragraph is None:
        raise SignatureGenerationError("会议记录母版中未找到表决结论段落。")
    result_source = deepcopy(result_paragraph._p)
    result_text = result_paragraph.text
    child_heading_paragraph = next((p for p in zone if "子议案如下" in p.text), zone[0])
    child_heading_source = deepcopy(child_heading_paragraph._p)
    child_paragraph = next(
        (
            zone[index + 1]
            for index, paragraph in enumerate(zone[:-1])
            if "子议案如下" in paragraph.text
        ),
        zone[0],
    )
    child_source = deepcopy(child_paragraph._p)
    _strip_numbering(child_source)
    boundary = paragraphs[end]
    # 最新规则的会议记录示例统一使用第 4 节点分编号（4.1、4.2.1）。
    numbered_as_four = True
    for paragraph in zone:
        _remove_paragraph(paragraph)

    for index, agenda in enumerate(meeting.agendas, 1):
        prefix = f"4.{index}" if numbered_as_four else f"{_cn_index(index)}、"
        action = "逐项审议通过" if agenda.children else "审议通过"
        separator = " " if numbered_as_four else ""
        _clone_before(boundary, title_source, f"{prefix}{separator}{action}《{agenda.title}》")
        conclusion = (
            "本议案按会议清单标注回避表决。"
            if agenda.recusal
            else re.sub(r"《[^》]*【[^】]*】[^》]*》", f"《{agenda.title}》", result_text)
        )
        _clone_before(boundary, result_source, conclusion)
        if agenda.children:
            _clone_before(boundary, child_heading_source, "子议案如下：")
            for child_index, child in enumerate(agenda.children, 1):
                child_prefix = f"{prefix}.{child_index}" if numbered_as_four else f"{index}.{child_index}"
                _clone_before(boundary, child_source, f"{child_prefix} {child}")


def _set_cell_text(cell: _Cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    _set_paragraph_text(paragraph, text)
    if _has_manual_number(text):
        _strip_numbering(paragraph._p)
    for extra in list(cell.paragraphs[1:]):
        cell._tc.remove(extra._p)


def _set_table_geometry(table: Table, column_widths: tuple[int, ...]) -> None:
    """Apply deterministic DXA geometry to a ballot table and every cell."""
    if len(table.columns) != len(column_widths):
        raise SignatureGenerationError("表决票表格列数与宽度规则不一致。")

    total_width = sum(column_widths)
    table.autofit = False
    table_properties = table._tbl.tblPr

    table_width = table_properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        table_properties.insert(0, table_width)
    table_width.set(qn("w:type"), "dxa")
    table_width.set(qn("w:w"), str(total_width))

    table_indent = table_properties.find(qn("w:tblInd"))
    if table_indent is None:
        table_indent = OxmlElement("w:tblInd")
        table_properties.append(table_indent)
    table_indent.set(qn("w:type"), "dxa")
    table_indent.set(qn("w:w"), str(BALLOT_TABLE_INDENT))

    table_layout = table_properties.find(qn("w:tblLayout"))
    if table_layout is None:
        table_layout = OxmlElement("w:tblLayout")
        table_properties.append(table_layout)
    table_layout.set(qn("w:type"), "fixed")

    for grid_column, width in zip(table._tbl.tblGrid.gridCol_lst, column_widths, strict=True):
        grid_column.set(qn("w:w"), str(width))

    for table_row in table._tbl.tr_lst:
        column_index = 0
        for table_cell in table_row.tc_lst:
            cell_properties = table_cell.get_or_add_tcPr()
            grid_span = cell_properties.find(qn("w:gridSpan"))
            span = int(grid_span.get(qn("w:val"))) if grid_span is not None else 1
            width = sum(column_widths[column_index:column_index + span])
            cell_width = cell_properties.find(qn("w:tcW"))
            if cell_width is None:
                cell_width = OxmlElement("w:tcW")
                cell_properties.insert(0, cell_width)
            cell_width.set(qn("w:type"), "dxa")
            cell_width.set(qn("w:w"), str(width))
            column_index += span


def _ballot_entries(agendas: tuple[AgendaInfo, ...]) -> list[tuple[str, bool]]:
    entries: list[tuple[str, bool]] = []
    for agenda in agendas:
        entries.append((f"《{agenda.title}》", agenda.recusal))
        entries.extend((f"    {index}. {child}", agenda.recusal) for index, child in enumerate(agenda.children, 1))
    return entries


def _is_ballot_table(table: Table) -> bool:
    if len(table.rows) < 3:
        return False
    header = "".join(cell.text for row in table.rows[:2] for cell in row.cells)
    compact = re.sub(r"\s+", "", header)
    return "表决意见" in compact or "表决情况" in compact


def _replace_ballot_tables(document: Document, meeting: MeetingInfo) -> None:
    ballot_tables = [table for table in document.tables if _is_ballot_table(table)]
    if not ballot_tables:
        raise SignatureGenerationError("表决票母版中未找到议案表决表格。")
    entries = _ballot_entries(meeting.agendas)
    for table in ballot_tables:
        template_row = deepcopy(table.rows[2]._tr)
        for row in list(table.rows)[2:]:
            table._tbl.remove(row._tr)
        for label, recusal in entries:
            table._tbl.append(deepcopy(template_row))
            row = table.rows[-1]
            _set_cell_text(row.cells[0], label)
            for cell in row.cells[1:]:
                _set_cell_text(cell, "")
            if recusal:
                merged = row.cells[1].merge(row.cells[3])
                _set_cell_text(merged, "回避表决")
                merged.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                paragraph = merged.paragraphs[0]
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
        if meeting.kind == "independent_directors":
            _set_table_geometry(table, INDEPENDENT_BALLOT_COLUMN_WIDTHS)


def _document_kind(template: Path) -> str:
    prefix = template.stem.lstrip().split(maxsplit=1)[0].rstrip("、.-")
    if prefix.startswith("01"):
        return "notice"
    if prefix.startswith("02"):
        return "ballot"
    if prefix.startswith("03"):
        return "resolution"
    if prefix.startswith("04"):
        return "record"
    raise SignatureGenerationError(f"无法从母版文件名识别文件类型：{template.name}")


def _output_name(document_kind: str, meeting: MeetingInfo) -> str:
    names = {
        "notice": f"01-关于召开{meeting.meeting_name}的通知.docx",
        "ballot": f"02-{meeting.meeting_name}表决办法及表决票.docx",
        "resolution": f"03-{meeting.meeting_name}决议.docx",
        "record": f"04-{meeting.meeting_name}会议记录.docx",
    }
    return names[document_kind]


def _process_meeting_template(
    template: Path,
    output_path: Path,
    package: ExcelMeetingPackage,
    meeting: MeetingInfo,
) -> None:
    document_kind = _document_kind(template)
    document = Document(template)
    if document_kind in {"notice", "ballot"}:
        _remove_agenda_guidance_paragraphs(document)
        _replace_agenda_line_placeholders(document, meeting)
    elif document_kind == "resolution":
        _replace_resolution_blocks(document, meeting)
    else:
        _replace_record_blocks(document, meeting)
    if document_kind == "ballot":
        _replace_ballot_tables(document, meeting)

    for paragraph in _all_paragraphs(document):
        replaced = _replace_metadata(paragraph.text, meeting, document_kind, package.memo_title)
        if replaced != paragraph.text:
            _set_paragraph_text(paragraph, replaced)
    document.save(output_path)

    rendered = Document(output_path)
    paragraphs = list(_all_paragraphs(rendered))
    text = "".join(paragraph.text for paragraph in paragraphs)
    if _placeholder_count(output_path):
        raise SignatureGenerationError(f"{template.name} 仍有未更新的 Excel 字段占位符。")
    missing = [agenda.title for agenda in meeting.agendas if agenda.title not in text]
    if missing:
        raise SignatureGenerationError(f"{template.name} 缺少议案：{'；'.join(missing[:3])}")
    overlapping = [
        paragraph.text.strip()
        for paragraph in paragraphs
        if _has_manual_number(paragraph.text)
        and _has_direct_numbering(paragraph)
        and ("审议" in paragraph.text or "《" in paragraph.text)
    ]
    if overlapping:
        raise SignatureGenerationError(f"{template.name} 存在自动编号与正文编号重叠：{overlapping[0]}")
    if document_kind in {"notice", "ballot"}:
        missing_wording = [
            agenda.title
            for agenda in meeting.agendas
            if agenda.children and f"逐项审议通过《{agenda.title}》" not in text
        ]
        if missing_wording:
            raise SignatureGenerationError(
                f"{template.name} 的子议案未使用‘逐项审议通过’表述：{missing_wording[0]}"
            )
        listed_children = [
            child
            for agenda in meeting.agendas
            for child in agenda.children
            if any(f"《{child}》" in paragraph.text for paragraph in rendered.paragraphs)
        ]
        if listed_children:
            raise SignatureGenerationError(
                f"{template.name} 的审议事项正文不应展开子议案：{listed_children[0]}"
            )
    if document_kind == "record":
        malformed = next(
            (
                paragraph.text.strip()
                for paragraph in paragraphs
                if re.match(r"\s*4\.\d+\s+\d+\b", paragraph.text)
            ),
            None,
        )
        if malformed:
            raise SignatureGenerationError(f"{template.name} 的子议案编号层级错误：{malformed}")
        uses_four_prefix = any(
            re.match(r"\s*4\.\d+\s+", paragraph.text) and "审议" in paragraph.text
            for paragraph in paragraphs
        )
        for agenda_index, agenda in enumerate(meeting.agendas, 1):
            for child_index, child in enumerate(agenda.children, 1):
                prefix = (
                    f"4.{agenda_index}.{child_index}"
                    if uses_four_prefix
                    else f"{agenda_index}.{child_index}"
                )
                if f"{prefix} {child}" not in text:
                    raise SignatureGenerationError(
                        f"{template.name} 缺少正确的子议案编号：{prefix} {child}"
                    )


def _create_job_dir(output_root: Path, package: ExcelMeetingPackage) -> tuple[str, Path]:
    output_root.mkdir(parents=True, exist_ok=True)
    base = f"{datetime.now():%Y%m%d-%H%M%S}-{package.meetings['board'].meeting_name}"
    for attempt in range(1, 1000):
        job_name = base if attempt == 1 else f"{base}-{attempt}"
        job_dir = output_root / job_name
        try:
            job_dir.mkdir(parents=True, exist_ok=False)
            return job_name, job_dir
        except FileExistsError:
            continue
    raise SignatureGenerationError("同一会议的生成任务过多，请稍后重试。")


def generate_complete_meeting_package(
    excel_path: Path,
    signature_template_root: Path,
    meeting_template_root: Path,
    output_root: Path,
) -> tuple[ExcelMeetingPackage, Path, list[dict[str, object]]]:
    package = parse_board_excel(excel_path)
    signature_templates = sorted(signature_template_root.glob("*.docx"))
    if not signature_templates:
        raise SignatureGenerationError("签字页模板库为空，请先恢复内置资料。")
    job_name, job_dir = _create_job_dir(output_root, package)
    documents: list[dict[str, object]] = []
    try:
        signature_dir = job_dir / "00-签字页"
        signature_dir.mkdir()
        for template in signature_templates:
            output_name = f"{package.meetings['board'].meeting_name}-{template.stem}.docx"
            output_path = signature_dir / output_name
            changed_parts = _rewrite_docx(template, output_path, package)
            if changed_parts == 0 or _placeholder_count(output_path):
                raise SignatureGenerationError(f"{template.name} 的会议字段未完整更新。")
            documents.append(
                {
                    "category": "signature",
                    "meeting_kind": "all",
                    "template": template.name,
                    "filename": str(output_path.relative_to(job_dir)),
                }
            )

        for meeting_kind in MEETING_KIND_ORDER:
            meeting = package.meetings.get(meeting_kind)
            if meeting is None:
                continue
            template_dir = meeting_template_root / meeting_kind
            templates = sorted(template_dir.rglob("*.docx"))
            if len(templates) != 4:
                raise SignatureGenerationError(
                    f"{MEETING_FOLDER_LABELS[meeting_kind]} 应有 4 份母版，当前找到 {len(templates)} 份。"
                )
            meeting_dir = job_dir / MEETING_FOLDER_LABELS[meeting_kind]
            meeting_dir.mkdir()
            for template in templates:
                kind = _document_kind(template)
                output_path = meeting_dir / _output_name(kind, meeting)
                try:
                    _process_meeting_template(template, output_path, package, meeting)
                except SignatureGenerationError as exc:
                    raise SignatureGenerationError(f"{template.name}：{exc}") from exc
                documents.append(
                    {
                        "category": kind,
                        "meeting_kind": meeting_kind,
                        "template": template.name,
                        "filename": str(output_path.relative_to(job_dir)),
                    }
                )

        return package, job_dir, documents
    except Exception:
        shutil.rmtree(job_dir, ignore_errors=True)
        raise
