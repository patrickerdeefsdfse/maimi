"""Read the supported employee workbook without evaluating user-supplied cells."""

from __future__ import annotations

import io
import re
import zipfile
from typing import Any

from openpyxl import load_workbook


MAX_WORKBOOK_BYTES = 10 * 1024 * 1024
MAX_EXPANDED_BYTES = 50 * 1024 * 1024
MAX_SHEET_ROWS = 10_000
MAX_SHEET_COLUMNS = 64
REQUIRED_HEADERS = ("工号", "姓名", "英文名")
EMPLOYEE_NUMBER = re.compile(r"M[0-9]{8}\Z")
CONTROL_CHARACTERS = re.compile(r"[\x00-\x1f\x7f]")


class ImportValidationError(ValueError):
    """Errors may be displayed directly; source cell contents are never echoed."""

    def __init__(self, errors: list[dict[str, Any]]):
        self.errors = errors
        super().__init__("名单校验未通过，请按提示修改 Excel 后重新上传。")


def _error(row: int | None, field: str, message: str) -> dict[str, Any]:
    return {"row": row, "field": field, "message": message}


def _fail(message: str) -> None:
    raise ImportValidationError([_error(None, "文件", message)])


def _empty(value: Any) -> bool:
    return value is None or isinstance(value, str) and not value.strip()


def _check_archive(data: bytes) -> None:
    if not data:
        _fail("文件为空，请上传标准 .xlsx 名单。")
    if len(data) > MAX_WORKBOOK_BYTES:
        _fail("Excel 文件不能超过 10 MB。")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 1000 or sum(item.file_size for item in entries) > MAX_EXPANDED_BYTES:
                _fail("Excel 文件展开后过大，请使用空白模板重新填写。")
            names = {entry.filename for entry in entries}
            if "xl/workbook.xml" not in names or "[Content_Types].xml" not in names:
                _fail("文件不是有效的 .xlsx 工作簿。")
            if any(name.lower().endswith("vbaproject.bin") for name in names):
                _fail("不支持包含宏的工作簿，请另存为标准 .xlsx 文件。")
            for entry in entries:
                if entry.filename.lower().endswith(".xml"):
                    # XML may use UTF-16/32; remove NUL code units before
                    # checking declarations so encoding cannot bypass this gate.
                    content = archive.read(entry).replace(b"\x00", b"")
                    if b"<!DOCTYPE" in content.upper() or b"<!ENTITY" in content.upper():
                        _fail("工作簿包含不支持的 XML 内容，请使用空白模板。")
    except ImportValidationError:
        raise
    except (zipfile.BadZipFile, RuntimeError, OSError, ValueError, NotImplementedError):
        _fail("无法读取文件。请确认它是未加密的标准 .xlsx 工作簿。")


def parse_workbook(data: bytes, max_rows: int = 500) -> list[dict]:
    """Return validated employees, preserving text IDs and original row numbers.

    The first row containing all three required headers is the header. Blank
    sheets and valid header-only sheets are ignored when selecting a data sheet.
    A populated sheet lacking the supported header is rejected rather than
    silently dropping employees from a workbook.
    """
    if not isinstance(max_rows, int) or isinstance(max_rows, bool) or not 1 <= max_rows <= MAX_SHEET_ROWS:
        raise ValueError("max_rows must be between 1 and 10000")
    _check_archive(data)
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except Exception:
        _fail("无法解析 Excel。请确认文件未加密，并使用下载的 .xlsx 模板。")

    errors: list[dict] = []
    candidate_tables: list[tuple[list, dict[str, int]]] = []
    try:
        if len(workbook.worksheets) > 20:
            _fail("工作表数量过多，请只保留需要导入的名单。")
        for sheet in workbook.worksheets:
            if (sheet.max_row or 0) > MAX_SHEET_ROWS or (sheet.max_column or 0) > MAX_SHEET_COLUMNS:
                errors.append(_error(None, "工作表", "表格范围过大，请使用模板并删除多余行列或格式。"))
                continue
            # Some workbook exporters declare an incorrect used range. Trusting
            # it would silently omit employee rows. Bound the actual traversal.
            sheet.reset_dimensions()
            rows = []
            range_error = False
            for row_number, row in enumerate(sheet.iter_rows(max_col=MAX_SHEET_COLUMNS + 1), start=1):
                if row_number > MAX_SHEET_ROWS or not _empty(row[-1].value):
                    errors.append(_error(None, "工作表", "表格范围过大，请使用模板并删除多余行列或格式。"))
                    range_error = True
                    break
                if any(not _empty(cell.value) for cell in row):
                    rows.append(row)
            if range_error:
                continue
            if not rows:
                continue
            header_position = None
            header_map: dict[str, int] = {}
            for index, row in enumerate(rows):
                labels = [cell.value.strip() if isinstance(cell.value, str) else "" for cell in row]
                if all(label in labels for label in REQUIRED_HEADERS):
                    header_position = index
                    for label in (*REQUIRED_HEADERS, "照片"):
                        if labels.count(label) > 1:
                            header_row = next(cell.row for cell in row if not _empty(cell.value))
                            errors.append(_error(header_row, label, "表头重复，请每个字段只保留一列。"))
                        if label in labels:
                            header_map[label] = labels.index(label)
                    break
            if header_position is None:
                errors.append(_error(None, "表头", "非空工作表必须包含“工号、姓名、英文名”三列表头。"))
                continue
            table_rows = rows[header_position + 1 :]
            if table_rows:
                candidate_tables.append((table_rows, header_map))
        if len(candidate_tables) > 1:
            errors.append(_error(None, "工作表", "发现多个包含员工数据的工作表，请每次只导入一个名单。"))
        if not candidate_tables:
            errors.append(_error(None, "名单", "没有找到员工数据，请在模板表头下填写名单。"))
        if errors:
            raise ImportValidationError(errors)

        rows, header_map = candidate_tables[0]
        if len(rows) > max_rows:
            _fail(f"每批最多导入 {max_rows} 人，请将名单拆分后上传。")
        employees: list[dict] = []
        seen: dict[str, int] = {}
        for cells in rows:
            row_number = next(cell.row for cell in cells if not _empty(cell.value))
            row_errors: list[dict] = []
            for cell in cells:
                if cell.data_type == "f":
                    row_errors.append(_error(row_number, "单元格", "不支持公式，请复制后粘贴为纯文本或值。"))
                    break
                if cell.data_type == "e":
                    row_errors.append(_error(row_number, "单元格", "单元格包含 Excel 错误，请修改后重试。"))
                    break
            values: dict[str, str] = {}
            for header in REQUIRED_HEADERS:
                value = cells[header_map[header]].value
                if _empty(value):
                    row_errors.append(_error(row_number, header, "必填项不能为空。"))
                    continue
                if not isinstance(value, str):
                    message = "工号必须为文本，以 M 开头并包含 8 位数字，不能使用数值格式。" if header == "工号" else "请填写文本内容。"
                    row_errors.append(_error(row_number, header, message))
                    continue
                value = value.strip()
                values[header] = value
                if CONTROL_CHARACTERS.search(value):
                    row_errors.append(_error(row_number, header, "内容不能包含换行、制表符或控制字符。"))
                if header == "工号":
                    if not EMPLOYEE_NUMBER.fullmatch(value):
                        row_errors.append(_error(row_number, header, "目前支持的工号格式为大写 M 加 8 位数字，例如 M00096389。"))
                    elif value in seen:
                        row_errors.append(_error(row_number, header, f"工号重复，与第 {seen[value]} 行相同。"))
                    else:
                        seen[value] = row_number
                elif len(value) > (24 if header == "姓名" else 64):
                    limit = 24 if header == "姓名" else 64
                    row_errors.append(_error(row_number, header, f"内容过长，最多允许 {limit} 个字符，请核对姓名。"))
            if "照片" in header_map and not _empty(cells[header_map["照片"]].value):
                row_errors.append(_error(row_number, "照片", "照片列请留空，系统将自动取照；缺失照片可在网页补传。请勿填写文件路径。"))
            errors.extend(row_errors)
            if not row_errors:
                employees.append({"row": row_number, "employee_no": values["工号"], "name": values["姓名"], "english_name": values["英文名"]})
        if errors:
            raise ImportValidationError(errors)
        return employees
    except ImportValidationError:
        raise
    except Exception:
        _fail("无法完整读取工作簿，请用空白模板重新填写并上传。")
    finally:
        workbook.close()
