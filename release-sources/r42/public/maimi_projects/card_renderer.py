"""Deterministic artwork and print documents for the supplied MEGMEET template."""

from __future__ import annotations

import io
import math
import os
import re
import unicodedata
import uuid
import zipfile
import shutil
from datetime import datetime, timezone
from functools import lru_cache
from pathlib import Path

import qrcode
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from fontTools.ttLib import TTFont
from PIL import Image, ImageDraw, ImageFont, ImageOps, UnidentifiedImageError

CANVAS_SIZE = (1063, 1673)
FONT_FILENAME = "SourceHanSansSC-Regular.otf"
MAX_PHOTO_BYTES = 10 * 1024 * 1024
MAX_PHOTO_PIXELS = 24_000_000
PHOTO_BOX = (214, 309, 843, 1126)
QR_BOX = (74, 1364, 294, 1582)
TEXT_COLOR = "#008f83"
PRINT_LAYOUT = {"card_width_mm": 85.0, "card_height_mm": 54.0,
                "gap_mm": 1.5, "page_margin_mm": 10.0}


class CardValidationError(ValueError):
    """An input or layout problem that can be shown to the operator safely."""


@lru_cache(maxsize=8)
def _font_characters(font_path: str) -> frozenset[int]:
    with TTFont(font_path, lazy=True) as font:
        return frozenset(font.getBestCmap())


@lru_cache(maxsize=128)
def _font(font_path: str, size: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(font_path, size=size)


def _validate_text(value: object, field: str, font_path: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise CardValidationError(f"{field}不能为空。")
    text = value.strip()
    if len(text) > 100 or any(unicodedata.category(c).startswith("C") for c in text):
        raise CardValidationError(f"{field}过长或包含不可打印字符，请修改后重试。")
    supported = _font_characters(font_path)
    missing = sorted({c for c in text if ord(c) not in supported})
    if missing:
        raise CardValidationError(f"{field}包含当前字库不支持的字符：{'、'.join(missing[:6])}。请联系管理员补充字库。")
    return text


def _draw_fitted_text(
    draw: ImageDraw.ImageDraw,
    text: str,
    field: str,
    box: tuple[int, int, int, int],
    font_path: str,
    max_size: int,
    min_size: int,
    *,
    left_aligned: bool = False,
    stroke: int = 0,
) -> None:
    """Shrink to a readable minimum; reject instead of clipping or truncating."""
    left, top, right, bottom = box
    for size in range(max_size, min_size - 1, -1):
        font = _font(font_path, size)
        bounds = draw.textbbox((0, 0), text, font=font, stroke_width=stroke)
        width, height = bounds[2] - bounds[0], bounds[3] - bounds[1]
        if width <= right - left and height <= bottom - top:
            x = left if left_aligned else left + (right - left - width) / 2
            y = top + (bottom - top - height) / 2
            draw.text(
                (round(x - bounds[0]), round(y - bounds[1])), text,
                font=font, fill=TEXT_COLOR, stroke_width=stroke,
                stroke_fill=TEXT_COLOR,
            )
            return
    raise CardValidationError(f"{field}太长，无法在工卡上清晰显示，请缩短后重试。")


def _read_photo(data: bytes) -> Image.Image:
    if not isinstance(data, bytes) or not data or len(data) > MAX_PHOTO_BYTES:
        raise CardValidationError("照片为空或超过 10 MB，请更换照片。")
    try:
        with Image.open(io.BytesIO(data)) as source:
            if source.width * source.height > MAX_PHOTO_PIXELS:
                raise CardValidationError("照片像素过大，请使用不超过 2400 万像素的照片。")
            source.load()
            photo = ImageOps.exif_transpose(source).convert("RGBA")
        background = Image.new("RGB", photo.size, "white")
        background.paste(photo, mask=photo.getchannel("A"))
        return background
    except (UnidentifiedImageError, OSError, ValueError, Image.DecompressionBombError) as exc:
        if isinstance(exc, CardValidationError):
            raise
        raise CardValidationError("照片不是有效图片，请更换后重试。") from None


def _atomic_save_image(card: Image.Image, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = output_path.with_name(f".{output_path.name}.{uuid.uuid4().hex}.tmp")
    try:
        card.save(temporary, format="PNG", dpi=(500, 500), optimize=True)
        os.replace(temporary, output_path)
    finally:
        temporary.unlink(missing_ok=True)


def _compose_card(employee: dict, photo_bytes: bytes, assets_dir: Path) -> Image.Image:
    """Compose a lossless, 1063 × 1673 pixel card using bundled artwork/fonts."""
    assets_dir = Path(assets_dir)
    font_path = str(assets_dir / "fonts" / FONT_FILENAME)
    if not Path(font_path).is_file():
        raise CardValidationError("工卡字库缺失，请联系管理员检查安装文件。")
    employee_no = _validate_text(employee.get("employee_no"), "工号", font_path)
    if not re.fullmatch(r"M\d{8}", employee_no, flags=re.ASCII):
        raise CardValidationError("工号必须是大写 M 加 8 位数字。")
    name = _validate_text(employee.get("name"), "姓名", font_path)
    english_name = _validate_text(employee.get("english_name"), "英文名", font_path)
    try:
        with Image.open(assets_dir / "card-background.bmp") as background:
            if background.size != CANVAS_SIZE:
                raise CardValidationError("工卡底图尺寸不正确，请联系管理员检查安装文件。")
            card = background.convert("RGB")
    except (OSError, UnidentifiedImageError):
        raise CardValidationError("工卡底图无法读取，请联系管理员检查安装文件。") from None

    photo = _read_photo(photo_bytes)
    x0, y0, x1, y1 = PHOTO_BOX
    photo_canvas = Image.new("RGB", (x1 - x0, y1 - y0), "white")
    fitted = ImageOps.contain(photo, photo_canvas.size, method=Image.Resampling.LANCZOS)
    photo_canvas.paste(fitted, ((photo_canvas.width - fitted.width) // 2, (photo_canvas.height - fitted.height) // 2))
    mask = Image.new("L", photo_canvas.size, 0)
    ImageDraw.Draw(mask).rounded_rectangle((0, 0, mask.width - 1, mask.height - 1), radius=54, fill=255)
    card.paste(photo_canvas, (x0, y0), mask)

    draw = ImageDraw.Draw(card)
    _draw_fitted_text(draw, name, "姓名", (258, 1238, 1030, 1312), font_path, 76, 44, stroke=1)
    _draw_fitted_text(draw, english_name, "英文名", (416, 1373, 1028, 1450), font_path, 68, 42, left_aligned=True)
    _draw_fitted_text(draw, employee_no, "工号", (560, 1512, 1028, 1581), font_path, 68, 48)

    # Keep the original M-prefixed identifier. The photo API's personPin is unrelated.
    qr = qrcode.QRCode(error_correction=qrcode.constants.ERROR_CORRECT_M, border=4)
    qr.add_data(employee_no)
    qr.make(fit=True)
    x0, y0, x1, y1 = QR_BOX
    module_count = qr.modules_count + 2 * qr.border
    qr.box_size = min(x1 - x0, y1 - y0) // module_count
    qr_image = qr.make_image(fill_color="black", back_color="white").convert("RGB")
    draw.rectangle((x0, y0, x1 - 1, y1 - 1), fill="white")
    card.paste(qr_image, (x0 + (x1 - x0 - qr_image.width) // 2, y0 + (y1 - y0 - qr_image.height) // 2))
    return card


def render_card(employee: dict, photo_bytes: bytes, output_path: Path, assets_dir: Path) -> None:
    """Retained for compatibility; printing does not persist preview artifacts."""
    with _compose_card(employee, photo_bytes, assets_dir) as card:
        _atomic_save_image(card, Path(output_path))


def render_print_card(employee: dict, photo_bytes: bytes, output_path: Path, assets_dir: Path) -> None:
    """Rotate original 500 dpi artwork losslessly, without a preview round trip."""
    with _compose_card(employee, photo_bytes, assets_dir) as card:
        with card.transpose(Image.Transpose.ROTATE_270) as rotated:
            _atomic_save_image(rotated, Path(output_path))


def validate_card(employee: dict, photo_bytes: bytes, assets_dir: Path) -> None:
    """Validate edits before committing them, without publishing any picture."""
    with _compose_card(employee, photo_bytes, assets_dir):
        pass


def _measurement(value: object, name: str, *, positive: bool) -> float:
    if isinstance(value, bool):
        raise CardValidationError(f"{name}必须是有效数字。")
    try:
        number = float(value)
    except (TypeError, ValueError):
        raise CardValidationError(f"{name}必须是有效数字。") from None
    if not math.isfinite(number) or (number <= 0 if positive else number < 0):
        condition = "大于 0" if positive else "不小于 0"
        raise CardValidationError(f"{name}必须是{condition}的有限数值。")
    return number


def _anchor_picture(paragraph, image_path: Path | io.BytesIO, x_mm: float, y_mm: float, width_mm: float, height_mm: float, number: int) -> None:
    """Anchor to the page, avoiding inline-image baseline and table-row overflow."""
    descriptor = str(image_path) if isinstance(image_path, Path) else image_path
    shape = paragraph.add_run().add_picture(descriptor, width=Mm(width_mm), height=Mm(height_mm))
    inline = shape._inline
    anchor = OxmlElement("wp:anchor")
    for key, value in {
        "distT": "0", "distB": "0", "distL": "0", "distR": "0",
        "simplePos": "0", "relativeHeight": str(number), "behindDoc": "0",
        "locked": "1", "layoutInCell": "0", "allowOverlap": "0",
    }.items():
        anchor.set(key, value)
    simple = OxmlElement("wp:simplePos")
    simple.set("x", "0")
    simple.set("y", "0")
    anchor.append(simple)
    for axis, offset in (("H", x_mm), ("V", y_mm)):
        position = OxmlElement(f"wp:position{axis}")
        position.set("relativeFrom", "page")
        coordinate = OxmlElement("wp:posOffset")
        coordinate.text = str(Mm(offset))
        position.append(coordinate)
        anchor.append(position)
    anchor.append(inline.find(qn("wp:extent")))
    extent = OxmlElement("wp:effectExtent")
    for side in ("l", "t", "r", "b"):
        extent.set(side, "0")
    anchor.append(extent)
    anchor.append(OxmlElement("wp:wrapNone"))
    properties = inline.find(qn("wp:docPr"))
    properties.set("descr", f"第 {number} 张员工工卡")
    anchor.append(properties)
    frame_properties = inline.find(qn("wp:cNvGraphicFramePr"))
    if frame_properties is not None:
        anchor.append(frame_properties)
    anchor.append(inline.find(qn("a:graphic")))
    inline.getparent().replace(inline, anchor)


def create_document(
    cards: list[Path],
    output_path: Path,
    card_width_mm: float = 85.0,
    card_height_mm: float = 54.0,
    gap_mm: float = 1.5,
    page_margin_mm: float = 10.0,
    *, prepared: bool = False,
) -> dict:
    """Print ten cards per A4 page, losslessly rotated to 85 x 54 mm.

    The individual downloadable PNG stays portrait (1063 x 1673 pixels).
    Only a print-local memory copy rotates clockwise; card artwork, QR modules,
    source files and employee order are unchanged. This is a fixed production
    print profile, not a model-selected fit-to-page or arbitrary resize.
    """
    width = _measurement(card_width_mm, "工卡宽度", positive=True)
    height = _measurement(card_height_mm, "工卡高度", positive=True)
    gap = _measurement(gap_mm, "工卡间距", positive=False)
    margin = _measurement(page_margin_mm, "页面边距", positive=False)
    actual = dict(zip(PRINT_LAYOUT, (width, height, gap, margin)))
    if any(abs(actual[key] - expected) > 1e-9 for key, expected in PRINT_LAYOUT.items()):
        raise CardValidationError("打印固定为A4每页10张，每张8.5×5.4厘米；请使用默认打印排版。")
    if not cards:
        raise CardValidationError("请至少选择一名照片完整的员工。")
    paths = [Path(card) for card in cards]
    for path in paths:
        try:
            with Image.open(path) as artwork:
                if artwork.size != (tuple(reversed(CANVAS_SIZE)) if prepared else CANVAS_SIZE):
                    raise CardValidationError("工卡成品图片尺寸不正确，请重新生成工卡后重试。")
                artwork.verify()
        except (OSError, ValueError, UnidentifiedImageError):
            raise CardValidationError("部分工卡源照片无法处理，请补传照片后重试。") from None
    columns, rows = 2, 5
    per_page = columns * rows
    pages = math.ceil(len(paths) / per_page)

    document = Document()
    document.core_properties.title = "员工工卡打印"
    document.core_properties.author = "MEGMEET"
    # Stable document bytes let authenticated retries reuse existing artifacts.
    document.core_properties.created = document.core_properties.modified = datetime(2020, 1, 1, tzinfo=timezone.utc)
    document.core_properties.subject = "A4每页10张，每张8.5×5.4厘米，请按实际大小（100%）打印"
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Mm(margin)
    section.header_distance = section.footer_distance = Mm(0)
    normal = document.styles["Normal"]
    normal.font.size = Pt(1)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.line_spacing = Pt(1)
    settings = document.settings.element
    if settings.find(qn("w:doNotAutoCompressPictures")) is None:
        settings.append(OxmlElement("w:doNotAutoCompressPictures"))

    # Center a full 2 x 5 grid, including the final partially occupied page.
    # Bounds are x=19.25..190.75 mm and y=10.5..286.5 mm on A4.
    grid_width = columns * width + (columns - 1) * gap
    grid_height = rows * height + (rows - 1) * gap
    x_start = (210 - grid_width) / 2
    y_start = (297 - grid_height) / 2
    for page_number in range(pages):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.page_break_before = page_number > 0
        paragraph.paragraph_format.keep_with_next = False
        paragraph.paragraph_format.widow_control = False
        for slot, card_path in enumerate(paths[page_number * per_page : (page_number + 1) * per_page]):
            row, col = divmod(slot, columns)
            rotated_png = card_path
            if not prepared:
                rotated_png = io.BytesIO()
                with Image.open(card_path) as source:
                    with source.transpose(Image.Transpose.ROTATE_270) as rotated:
                        rotated.save(rotated_png, format="PNG", dpi=(500, 500))
                rotated_png.seek(0)
            _anchor_picture(
                paragraph, rotated_png,
                x_start + col * (width + gap), y_start + row * (height + gap),
                width, height, page_number * per_page + slot + 1,
            )
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = output_path.with_name(f".{output_path.name}.{uuid.uuid4().hex}.tmp")
    canonical = temporary.with_suffix('.canonical')
    try:
        document.save(temporary)
        with zipfile.ZipFile(temporary) as source, zipfile.ZipFile(canonical, 'w') as target:
            for entry in source.infolist():
                fixed = zipfile.ZipInfo(entry.filename, (2020, 1, 1, 0, 0, 0))
                fixed.compress_type = entry.compress_type
                fixed.external_attr = entry.external_attr
                with source.open(entry) as inp, target.open(fixed, 'w') as out:
                    shutil.copyfileobj(inp, out, 1024 * 1024)
        os.replace(canonical, output_path)
    finally:
        temporary.unlink(missing_ok=True)
        canonical.unlink(missing_ok=True)
    return {"pages": pages, "columns": columns, "rows": rows, "count": len(paths)}


class CardPrintSizeError(ValueError):
    """The single bounded output cannot fit, without changing image quality."""


class _BoundedPrintBuffer(io.BytesIO):
    def __init__(self, maximum):
        super().__init__()
        self.maximum = maximum

    def write(self, value):
        if self.tell() + len(value) > self.maximum:
            raise CardPrintSizeError('single print file exceeds limit')
        return super().write(value)


def render_print_bytes(employee: dict, photo_bytes: bytes, assets_dir: Path) -> bytes:
    """One transient, lossless 500 dpi card; never writes to /tmp or a preview."""
    with _compose_card(employee, photo_bytes, assets_dir) as card:
        with card.transpose(Image.Transpose.ROTATE_270) as rotated:
            out = io.BytesIO()
            rotated.save(out, format='PNG', dpi=(500, 500), optimize=True)
            return out.getvalue()


class SinglePrintDocument:
    """Streaming OPC archive with one artwork at a time, bounded by file bytes.

    python-docx is used only for the small layout/anchor skeleton. Actual image
    payloads are written directly into the ZIP and never retained by image_parts.
    No employee PNG or completed DOCX is written to the gateway's 32 MiB /tmp.
    """
    def __init__(self, count: int, maximum: int, **layout):
        if type(count) is not int or count <= 0:
            raise CardValidationError('请至少选择一名照片完整的员工。')
        values = {**PRINT_LAYOUT, **layout}
        for key, expected in PRINT_LAYOUT.items():
            if abs(_measurement(values[key], key, positive=key.startswith('card_')) - expected) > 1e-9:
                raise CardValidationError('打印固定为A4每页10张，每张8.5×5.4厘米；请使用默认打印排版。')
        self.count, self.added = count, 0
        self.buffer = _BoundedPrintBuffer(maximum)
        self.archive = zipfile.ZipFile(self.buffer, 'w', compression=zipfile.ZIP_DEFLATED)
        self.document = Document()
        d = self.document
        d.core_properties.title = '员工工卡打印'
        d.core_properties.author = 'MEGMEET'
        d.core_properties.created = d.core_properties.modified = datetime(2020, 1, 1, tzinfo=timezone.utc)
        d.core_properties.subject = 'A4每页10张，每张8.5×5.4厘米，请按实际大小（100%）打印'
        section = d.sections[0]
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Mm(10)
        section.header_distance = section.footer_distance = Mm(0)
        normal = d.styles['Normal']
        normal.font.size = Pt(1)
        normal.paragraph_format.space_after = normal.paragraph_format.space_before = Pt(0)
        normal.paragraph_format.line_spacing = Pt(1)
        if d.settings.element.find(qn('w:doNotAutoCompressPictures')) is None:
            d.settings.element.append(OxmlElement('w:doNotAutoCompressPictures'))
        tiny = io.BytesIO()
        Image.new('RGB', (1, 1), 'white').save(tiny, 'PNG')
        self.tiny = tiny.getvalue()
        self.relationships = []
        self.paragraph = None

    def _write(self, name, value, *, image=False):
        entry = zipfile.ZipInfo(name, (2020, 1, 1, 0, 0, 0))
        entry.compress_type = zipfile.ZIP_STORED if image else zipfile.ZIP_DEFLATED
        self.archive.writestr(entry, value)

    def add(self, artwork: bytes):
        if self.added >= self.count:
            raise CardValidationError('工卡数量与名单不一致。')
        with Image.open(io.BytesIO(artwork)) as image:
            if image.size != (1673, 1063) or image.format != 'PNG':
                raise CardValidationError('工卡成品图片尺寸不正确。')
            image.verify()
        slot = self.added % 10
        if slot == 0:
            self.paragraph = self.document.add_paragraph()
            self.paragraph.paragraph_format.page_break_before = self.added > 0
            self.paragraph.paragraph_format.keep_with_next = False
            self.paragraph.paragraph_format.widow_control = False
        number = self.added + 1
        row, col = divmod(slot, 2)
        _anchor_picture(self.paragraph, io.BytesIO(self.tiny), 19.25 + col * 86.5,
                        10.5 + row * 55.5, 85, 54, number)
        anchor = self.paragraph._p.findall('.//' + qn('wp:anchor'))[-1]
        rid = 'megmeet_card_' + str(number)
        anchor.find('.//' + qn('a:blip')).set(qn('r:embed'), rid)
        name = f'media/card-{number:04}.png'
        self._write('word/' + name, artwork, image=True)
        self.relationships.append((rid, name))
        self.added += 1

    def finish(self) -> bytes:
        # Preserve the OPC default namespace. Some Office importers reject an
        # ns0-prefixed Relationships root although Python can reopen it.
        from lxml import etree as ET
        if self.added != self.count:
            raise CardValidationError('工卡数量与名单不一致，未生成部分文件。')
        skeleton = io.BytesIO()
        self.document.save(skeleton)
        relation_ns = 'http://schemas.openxmlformats.org/package/2006/relationships'
        with zipfile.ZipFile(skeleton) as template:
            relations = ET.fromstring(template.read('word/_rels/document.xml.rels'))
            for child in list(relations):
                if child.get('Type', '').endswith('/image'):
                    relations.remove(child)
            for rid, target in self.relationships:
                ET.SubElement(relations, '{' + relation_ns + '}Relationship',
                              {'Id': rid, 'Target': target, 'Type': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/image'})
            for entry in template.infolist():
                if entry.filename.startswith('word/media/'):
                    continue
                value = (ET.tostring(relations, encoding='UTF-8', xml_declaration=True, standalone=True)
                         if entry.filename == 'word/_rels/document.xml.rels' else template.read(entry))
                self._write(entry.filename, value)
        self.archive.close()
        return self.buffer.getvalue()

    def close(self):
        # On cancellation or an oversized batch, discard the incomplete ZIP.
        try:
            self.archive.close()
        except CardPrintSizeError:
            self.archive.fp = None
        self.buffer.close()
